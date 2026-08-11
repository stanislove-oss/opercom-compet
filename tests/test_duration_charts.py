"""Регрессия по duration-графикам слайдов 13-14.

История бага: у этих графиков в шаблоне ДВЕ оси значений. Основной
percentStacked-плот (серии 5/10/15/20/25/''/50) висит на первой оси,
а служебная серия Ср.Хроно — на второй, с намеренно «странным» масштабом
(например min=-10, max=20). Именно этот масштаб делает 100%-столбец
Ср.Хроно узкой красной капсулой ВНУТРИ столбца.

Код заливки данных выравнивал масштаб по ВСЕМ c:valAx (фикс комбо-графиков
слайдов 24/30). Для слайдов 13-14 это ломало картинку: вторичная ось
становилась 0..1, капсула Ср.Хроно растягивалась на всю высоту столбца —
подсвечивалась не нужная часть столбца, а столбец целиком.

Тесты ниже фиксируют оба условия: вторичную ось не трогаем, а сама капсула
остаётся видимой рамкой без заливки.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

pptx = pytest.importorskip("pptx")

from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402

from functions.build_presentations import (  # noqa: E402
    DURATION_CHART_SERIES_ORDER,
    find_shape_by_name,
    get_avg_duration_axis_ids,
    prepare_duration_chart_series,
    update_chart_data_by_shape_name,
)

TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent
    / "templates_pptx"
    / "Х5_Оперком_март_v3_fin_named.pptx"
)

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}

MONTHS_RU = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]

#: Графики длительностей из шаблона: (слайд, имя shape).
DURATION_CHARTS = [
    (13, "chart__pyaterochka_duration_nat"),
    (13, "chart__magnit_duration_nat"),
    (14, "chart__pyaterochka_duration_reg"),
    (14, "chart__magnit_duration_reg"),
]


def make_categories(count=19):
    labels = [f"{month}.25" for month in MONTHS_RU] + [f"{month}.26" for month in MONTHS_RU]
    return labels[:count]


def make_series(count=19):
    return {
        "5": [200 + 10 * i for i in range(count)],
        "10": [2000 + 30 * i for i in range(count)],
        "15": [300 + 5 * i for i in range(count)],
        "20": [150] * count,
        "25": [100] * count,
        "": [0] * count,
        "50": [0] * count,
        "Ср.Хроно": [9.5 + 0.2 * i for i in range(count)],
    }


def get_chart(presentation, slide_num, shape_name):
    shape = find_shape_by_name(presentation.slides[slide_num - 1], shape_name)
    assert shape is not None and shape.has_chart, f"нет графика {shape_name} на слайде {slide_num}"
    return shape.chart


def axis_scaling(chart, axis_id):
    for value_axis in chart._chartSpace.valAx_lst:
        ax_id = value_axis.find("c:axId", NS)
        if ax_id is None or ax_id.get("val") != axis_id:
            continue
        scaling = value_axis.find("c:scaling", NS)
        maximum = scaling.find("c:max", NS)
        minimum = scaling.find("c:min", NS)
        return (
            float(maximum.get("val")) if maximum is not None else None,
            float(minimum.get("val")) if minimum is not None else None,
        )
    return None


def series_names(chart):
    names = []
    plot_area = chart._chartSpace.find(".//c:plotArea", NS)
    for ser in plot_area.findall(".//c:ser", NS):
        value = ser.find("c:tx//c:v", NS)
        names.append(value.text if value is not None and value.text is not None else "")
    return names


def avg_duration_series_element(chart):
    plot_area = chart._chartSpace.find(".//c:plotArea", NS)
    for ser in plot_area.findall(".//c:ser", NS):
        value = ser.find("c:tx//c:v", NS)
        name = value.text if value is not None and value.text is not None else ""
        if str(name).strip().lower().replace("ё", "е") in {"ср.хроно", "ср. хроно"}:
            return ser
    return None


@pytest.fixture(scope="module")
def filled_presentation():
    if not TEMPLATE_PATH.exists():
        pytest.skip(f"нет шаблона {TEMPLATE_PATH}")

    categories = make_categories()
    series_data = make_series(len(categories))
    presentation = pptx.Presentation(str(TEMPLATE_PATH))

    for slide_num, shape_name in DURATION_CHARTS:
        update_chart_data_by_shape_name(
            presentation=presentation,
            slide=slide_num,
            shape_name=shape_name,
            chart_type=XL_CHART_TYPE.COLUMN_STACKED_100,
            categories=categories,
            series_data=series_data,
            color_dict={"Ср. хроно": (255, 0, 0)},
            size_label=7,
        )

    return presentation


@pytest.fixture(scope="module")
def template_presentation():
    if not TEMPLATE_PATH.exists():
        pytest.skip(f"нет шаблона {TEMPLATE_PATH}")
    return pptx.Presentation(str(TEMPLATE_PATH))


@pytest.mark.parametrize("slide_num,shape_name", DURATION_CHARTS)
def test_secondary_axis_of_avg_duration_plot_is_untouched(
    template_presentation, filled_presentation, slide_num, shape_name
):
    """Масштаб вторичной оси Ср.Хроно должен остаться шаблонным."""
    template_chart = get_chart(template_presentation, slide_num, shape_name)
    filled_chart = get_chart(filled_presentation, slide_num, shape_name)

    protected_axis_ids = get_avg_duration_axis_ids(template_chart)
    assert protected_axis_ids, "не нашли служебный плот Ср.Хроно в шаблоне"

    checked = 0
    for axis_id in protected_axis_ids:
        expected = axis_scaling(template_chart, axis_id)
        if expected is None:
            continue
        checked += 1
        assert axis_scaling(filled_chart, axis_id) == expected, (
            f"{shape_name}: масштаб вторичной оси {axis_id} переписан "
            f"({axis_scaling(filled_chart, axis_id)} вместо {expected}). "
            "Ср.Хроно растянется на весь столбец."
        )

    assert checked, "во вторичной оси шаблона нет max/min — тест нечего проверять"


@pytest.mark.parametrize("slide_num,shape_name", DURATION_CHARTS)
def test_primary_axis_is_rescaled_to_100_percent(filled_presentation, slide_num, shape_name):
    """Основную ось 100%-stacked всё так же приводим к 0..100%."""
    chart = get_chart(filled_presentation, slide_num, shape_name)
    protected_axis_ids = get_avg_duration_axis_ids(chart)

    primary_axes = [
        value_axis for value_axis in chart._chartSpace.valAx_lst
        if value_axis.find("c:axId", NS).get("val") not in protected_axis_ids
    ]
    assert primary_axes, "не нашли основную ось значений"

    for value_axis in primary_axes:
        scaling = value_axis.find("c:scaling", NS)
        assert float(scaling.find("c:max", NS).get("val")) == 1.0
        assert float(scaling.find("c:min", NS).get("val")) == 0.0


@pytest.mark.parametrize("slide_num,shape_name", DURATION_CHARTS)
def test_avg_duration_marker_is_visible_outline(filled_presentation, slide_num, shape_name):
    """Капсула Ср.Хроно — рамка без заливки, а не скрытая серия."""
    chart = get_chart(filled_presentation, slide_num, shape_name)
    ser = avg_duration_series_element(chart)
    assert ser is not None, f"{shape_name}: серия Ср.Хроно потерялась"

    spPr = ser.find("c:spPr", NS)
    assert spPr is not None
    assert spPr.find("a:noFill", NS) is not None, "у капсулы Ср.Хроно не должно быть заливки"

    line = spPr.find("a:ln", NS)
    assert line is not None
    assert line.find("a:noFill", NS) is None, "рамка капсулы не должна быть скрыта"

    color = line.find("a:solidFill/a:srgbClr", NS)
    assert color is not None and color.get("val").upper() == "FF0000"
    assert int(line.get("w") or 0) > 0, "нулевая толщина рамки = невидимая капсула"


@pytest.mark.parametrize("slide_num,shape_name", DURATION_CHARTS)
def test_series_set_and_order_match_template(
    template_presentation, filled_presentation, slide_num, shape_name
):
    """Набор и порядок серий комбо-графика менять нельзя."""
    template_chart = get_chart(template_presentation, slide_num, shape_name)
    filled_chart = get_chart(filled_presentation, slide_num, shape_name)

    assert series_names(filled_chart) == series_names(template_chart)


def test_prepare_duration_chart_series_keeps_template_order():
    prepared = prepare_duration_chart_series(
        {'5"': [1, 2], '12': [3, 4], '40': [5, 6], "Ср.Хроно": [7.0, 8.0]},
        categories=["янв.25", "фев.25"],
    )

    assert list(prepared) == DURATION_CHART_SERIES_ORDER
    # 12 секунд попадают в bucket 15, 40 — в bucket 50.
    assert prepared["5"] == [1, 2]
    assert prepared["15"] == [3, 4]
    assert prepared["50"] == [5, 6]
    assert prepared[""] == [0, 0]
    assert prepared["Ср.Хроно"] == [7.0, 8.0]
