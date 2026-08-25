from pptx import Presentation
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def extract_slide_tags(pptx_path):
    """Возвращает dict {тег: номер_слайда} из размеченной презентации."""
    prs = Presentation(pptx_path)
    result = {}
    for i, slide in enumerate(prs.slides, start=1):
        try:
            notes = slide.notes_slide.notes_text_frame.text
            for line in notes.splitlines():
                if line.strip().startswith("SLIDE_TAG:"):
                    tag = line.replace("SLIDE_TAG:", "").strip()
                    result[tag] = i
                    break
        except Exception:
            pass
    return result


def save_tags_to_excel(pptx_path: str, output_path: str) -> dict:
    tags = extract_slide_tags(pptx_path)  # использует функцию выше

    wb = Workbook()
    ws = wb.active
    ws.title = "Теги слайдов"

    header_fill = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    row_font    = Font(name="Arial", size=10)
    alt_fill    = PatternFill("solid", start_color="EBF3FB", end_color="EBF3FB")
    thin        = Side(style="thin", color="BFBFBF")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    center      = Alignment(horizontal="center", vertical="center")
    left        = Alignment(horizontal="left", vertical="center")

    for col, (h, w) in enumerate(zip(["№ слайда", "Тег"], [12, 45]), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22

    for i, (tag, slide_num) in enumerate(sorted(tags.items(), key=lambda x: x[1]), start=2):
        fill = alt_fill if i % 2 == 0 else None
        for col, (val, aln) in enumerate([(slide_num, center), (tag, left)], 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.font = row_font
            cell.alignment = aln
            cell.border = border
            if fill:
                cell.fill = fill
        ws.row_dimensions[i].height = 18

    ws.freeze_panes = "A2"
    wb.save(output_path)
    return tags
