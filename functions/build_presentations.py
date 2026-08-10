#Основная функция которая заливает данные в график

import math
from numbers import Number
import re
import io 
import copy
import openpyxl

from lxml import etree
import pandas as pd
from pptx.chart.axis import ValueAxis
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt
from datetime import date
from openpyxl.utils import column_index_from_string

MONTHS_RU = {
    'янв': 1, 'фев': 2, 'мар': 3, 'апр': 4, 'май': 5, 'июн': 6,
    'июл': 7, 'авг': 8, 'сен': 9, 'окт': 10, 'ноя': 11, 'дек': 12,
}
EXCEL_EPOCH = date(1899, 12, 30)


TAG_PATTERN = re.compile(r"\{([A-Za-zА-Яа-яЁё0-9_\-]+)\}")
NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}


def normalize_tag_key(key):
    """
    Приводит ключ тега к единому виду:
    "{first_month}" -> "first_month"
    "first_month" -> "first_month"
    """
    return str(key).strip().strip("{}")


def prepare_tags_dict(tags_dict):
    """
    Нормализует ключи словаря тегов.
    Значения None превращает в пустую строку.
    """
    result = {}

    for key, value in tags_dict.items():
        clean_key = normalize_tag_key(key)

        if value is None:
            value = ""

        result[clean_key] = str(value)

    return result


def get_slide_tags(tag_values, slide_num):
    """
    Возвращает итоговый словарь тегов для конкретного слайда.

    Поддерживает:
    1) плоский словарь:
       {"first_month": "Январь"}

    2) словарь с global/slides:
       {
           "global": {...},
           "slides": {
               2: {...},
               4: {...}
           }
       }
    """

    # Вариант global/slides
    if isinstance(tag_values, dict) and (
        "global" in tag_values or "slides" in tag_values
    ):
        result = {}

        global_tags = tag_values.get("global", {})
        slide_tags = tag_values.get("slides", {}).get(slide_num, {})

        # На случай если ключи слайдов строковые: "2", "4"
        if not slide_tags:
            slide_tags = tag_values.get("slides", {}).get(str(slide_num), {})

        result.update(prepare_tags_dict(global_tags))
        result.update(prepare_tags_dict(slide_tags))

        return result

    # Обычный плоский словарь
    return prepare_tags_dict(tag_values)


def replace_tags_in_text(text, tags):
    """
    Заменяет все теги вида {tag_name} в обычной строке.
    Если тега нет в словаре, оставляет его как есть.
    """

    def replace_match(match):
        tag_name = match.group(1)
        return tags.get(tag_name, match.group(0))

    return TAG_PATTERN.sub(replace_match, text)


def replace_tags_in_paragraph(paragraph, tags):
    """
    Заменяет теги внутри paragraph.

    Сначала пытается заменить по runs, чтобы сохранить форматирование.
    Если тег был разбит PowerPoint'ом на несколько runs,
    делает fallback: заменяет весь текст paragraph целиком.
    """

    if not paragraph.runs:
        return 0

    replaced_count = 0

    # 1. Обычная замена внутри каждого run
    for run in paragraph.runs:
        old_text = run.text
        new_text = replace_tags_in_text(old_text, tags)

        if new_text != old_text:
            run.text = new_text
            replaced_count += 1

    # Fallback for tags split across multiple runs. paragraph.text preserves soft line breaks (\x0b).
    full_text = paragraph.text
    new_full_text = replace_tags_in_text(full_text, tags)

    if new_full_text != full_text:
        # paragraph.text preserves soft line breaks, but recreates runs with default
        # formatting. Capture the first run style and reapply it after replacement.
        first_font = paragraph.runs[0].font
        font_size = first_font.size
        font_name = first_font.name
        font_bold = first_font.bold
        font_italic = first_font.italic
        font_underline = first_font.underline
        try:
            font_rgb = first_font.color.rgb
        except Exception:
            font_rgb = None

        paragraph.text = new_full_text

        for run in paragraph.runs:
            if font_size is not None:
                run.font.size = font_size
            if font_name is not None:
                run.font.name = font_name
            if font_bold is not None:
                run.font.bold = font_bold
            if font_italic is not None:
                run.font.italic = font_italic
            if font_underline is not None:
                run.font.underline = font_underline
            if font_rgb is not None:
                run.font.color.rgb = font_rgb
        replaced_count += 1

    return replaced_count


def replace_tags_in_text_frame(text_frame, tags):
    replaced_count = 0

    for paragraph in text_frame.paragraphs:
        replaced_count += replace_tags_in_paragraph(paragraph, tags)

    return replaced_count


def iter_all_shapes(shapes):
    """
    Рекурсивно проходит по shape'ам, включая grouped shapes.
    """
    for shape in shapes:
        yield shape

        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for inner_shape in iter_all_shapes(shape.shapes):
                yield inner_shape


def replace_ppt_tags(presentation, tag_values, include_notes=False):
    """
    Заменяет теги вида {tag_name} во всей презентации.

    presentation: объект Presentation
    tag_values: словарь тегов:
        либо плоский:
            {"first_month": "Январь"}

        либо с global/slides:
            {
                "global": {...},
                "slides": {
                    2: {...},
                    4: {...}
                }
            }

    include_notes: если True, также заменяет теги в заметках к слайдам.

    Возвращает количество заменённых текстовых блоков/runs.
    """

    total_replaced = 0

    for slide_idx, slide in enumerate(presentation.slides, start=1):
        slide_tags = get_slide_tags(tag_values, slide_idx)

        for shape in iter_all_shapes(slide.shapes):

            # Обычные текстовые блоки
            if shape.has_text_frame:
                total_replaced += replace_tags_in_text_frame(
                    shape.text_frame,
                    slide_tags
                )

            # Таблицы
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        total_replaced += replace_tags_in_text_frame(
                            cell.text_frame,
                            slide_tags
                        )

            # Заголовки chart'ов, если там вдруг тоже есть теги
            if shape.has_chart:
                chart = shape.chart

                if chart.has_title:
                    try:
                        total_replaced += replace_tags_in_text_frame(
                            chart.chart_title.text_frame,
                            slide_tags
                        )
                    except Exception:
                        pass

        # Заметки к слайдам
        if include_notes:
            try:
                total_replaced += replace_tags_in_text_frame(
                    slide.notes_slide.notes_text_frame,
                    slide_tags
                )
            except Exception:
                pass

    return total_replaced



# -----------------------------
# Text wrapping / margins cleanup
# -----------------------------

def _norm(name):
    return str(name).strip().lower().replace("ё", "е")



def _set_bodypr_nowrap_and_zero_margins(body_pr):
    """
    Настраивает XML bodyPr:
    - wrap='none' — запрет переноса
    - поля = 0

    В Open XML размеры inset'ов задаются в EMU, поэтому строка '0' корректна.
    """
    if body_pr is None:
        return 0

    body_pr.set('wrap', 'none')
    body_pr.set('lIns', '0')
    body_pr.set('rIns', '0')
    body_pr.set('tIns', '0')
    body_pr.set('bIns', '0')
    return 1

def _resize_range(formula, n_points):
    """
    'Лист1!$B$2:$B$16' + 18 точек -> 'Лист1!$B$2:$B$19'
    Меняет только номер последней строки.
    """
    m = re.match(r'^(.*\$)(\d+)(:\$[A-Z]+\$)(\d+)$', formula)
    if m is None:
        return formula                    # непонятный формат — не трогаем
    first_row = int(m.group(2))
    last_row = first_row + n_points - 1
    return f"{m.group(1)}{m.group(2)}{m.group(3)}{last_row}"

def _cat_to_excel_serial(label):
    """'янв.25' -> 45658. Возвращает None, если формат не распознан."""
    m = re.match(r'^([а-я]{3})\.?\s*(\d{2})$', str(label).strip().lower())
    if m is None:
        return None
    month = MONTHS_RU.get(m.group(1))
    if month is None:
        return None
    year = 2000 + int(m.group(2))
    return (date(year, month, 1) - EXCEL_EPOCH).days

def _columns_by_series(chart):
    """{'5': 'B', '10': 'C', ...} из формул c:val//c:f"""
    NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    out = {}
    pa = chart._chartSpace.find('.//c:plotArea', NS)
    for ser in pa.findall('.//c:ser', NS):
        v = ser.find('c:tx//c:v', NS)
        name = v.text if v is not None and v.text is not None else ""
        f = ser.find('c:val//c:f', NS).text
        m = re.search(r'\$([A-Z]+)\$\d+', f)
        if m:
            out[name] = m.group(1)
    return out

def _update_chart_workbook(chart, categories, series_values):
    """
    Переписывает встроенный Excel графика.
    categories: список подписей ('янв.25', ...) — конвертируются в серийные номера
    series_values: dict {имя_серии: [числа]}
    """
    cols = _columns_by_series(chart)
    cols = { _norm(k): v for k, v in cols.items()}
    

    xlsx_part = chart.part.chart_workbook.xlsx_part
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_part.blob))
    ws = wb.worksheets[0]
    ws["A1"] = "cat"
    for i, (name, col) in enumerate(cols.items(), start=1):
        ws[f"{col}1"] = name

    # 1. категории в колонку A, строки 2..(1+N)
    #    значение = _cat_to_excel_serial(label)
    #    number_format = 'mmm-yy'
    for i, label in enumerate(categories):
        row = i + 2
        serial = _cat_to_excel_serial(label)
        if serial is None:
            continue
        ws.cell(row=row, column=1, value=serial)
        ws.cell(row=row, column=1).number_format = "mmm-yy"

    # 2. значения серий в свои колонки, те же строки
    #    колонку берём из cols[name]
    #    ячейка: ws[f"{cols[name]}{row}"]
    for series, values in series_values.items():
        key = _norm(series)
        if key not in cols:
            continue
        col = cols[key]
        for i, val in enumerate(values):
            row = i + 2
            cell = ws.cell(row=row, column=column_index_from_string(col))
            if val is None:
                continue
            else:
                cell.value = val
                cell.number_format = 'General'

    # 3. удалить всё ниже последней строки
    last_row = 1 + len(categories)
    if ws.max_row > last_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_part.blob = buf.getvalue()
    print("  воркбук записан, строк:", ws.max_row)



def _ensure_chart_txpr_with_bodypr(chart_xml_element):
    """
    Для chart XML-элемента, например c:dLbls или c:dLbl,
    гарантирует наличие c:txPr/a:bodyPr и возвращает bodyPr.
    """
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    tx_pr = chart_xml_element.find(f'{{{ns_c}}}txPr')
    if tx_pr is None:
        tx_pr = etree.SubElement(chart_xml_element, f'{{{ns_c}}}txPr')

    body_pr = tx_pr.find(f'{{{ns_a}}}bodyPr')
    if body_pr is None:
        body_pr = etree.SubElement(tx_pr, f'{{{ns_a}}}bodyPr')

    lst_style = tx_pr.find(f'{{{ns_a}}}lstStyle')
    if lst_style is None:
        tx_pr.append(etree.Element(f'{{{ns_a}}}lstStyle'))

    # Для валидного txPr PowerPoint обычно ожидает хотя бы один a:p
    paragraph = tx_pr.find(f'{{{ns_a}}}p')
    if paragraph is None:
        tx_pr.append(etree.Element(f'{{{ns_a}}}p'))

    return body_pr


def configure_data_labels_no_wrap(data_labels):
    """
    Убирает перенос и поля у подписей данных графика на уровне XML.
    python-pptx не даёт полноценного API для word wrap/margins у DataLabels,
    поэтому меняем c:dLbls/c:txPr/a:bodyPr напрямую.
    """
    try:
        dLbls = data_labels._element
    except Exception:
        return 0

    changed = 0

    try:
        body_pr = _ensure_chart_txpr_with_bodypr(dLbls)
        changed += _set_bodypr_nowrap_and_zero_margins(body_pr)
    except Exception:
        pass

    # На случай индивидуальных подписей точек c:dLbl — тоже приводим их к тому же виду
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    try:
        for dLbl in dLbls.findall(f'{{{ns_c}}}dLbl'):
            body_pr = _ensure_chart_txpr_with_bodypr(dLbl)
            changed += _set_bodypr_nowrap_and_zero_margins(body_pr)
    except Exception:
        pass

    return changed


def reset_data_labels_xml(series):
    """Удаляет кастомный диапазон меток и индивидуальные метки точек"""
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    ns_c15 = 'http://schemas.microsoft.com/office/drawing/2012/chart'
    
    ser_elem = series._element
    
    # Удаляем datalabelsRange из extLst серии
    for ext_lst in ser_elem.findall(f'{{{ns_c}}}extLst'):
        for ext in ext_lst.findall(f'{{{ns_c}}}ext'):
            for dlbl_range in ext.findall(f'{{{ns_c15}}}datalabelsRange'):
                ext.remove(dlbl_range)
    
    # Удаляем showDataLabelsRange из dLbls
    for dLbls in ser_elem.findall(f'{{{ns_c}}}dLbls'):
        for ext_lst in dLbls.findall(f'{{{ns_c}}}extLst'):
            for ext in ext_lst.findall(f'{{{ns_c}}}ext'):
                for show in ext.findall(f'{{{ns_c15}}}showDataLabelsRange'):
                    ext.remove(show)
        # Удаляем индивидуальные метки точек
        for dLbl in dLbls.findall(f'{{{ns_c}}}dLbl'):
            dLbls.remove(dLbl)


def is_bright_color(rgb_color):
    """
    Проверяет, является ли цвет светлым (для выбора контрастного цвета текста)
    """
    r, g, b = rgb_color
    # Формула для вычисления яркости: 0.299*R + 0.587*G + 0.114*B
    brightness = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return brightness > 0.5


def is_stacked_chart(chart_type):
    return chart_type in [
        XL_CHART_TYPE.COLUMN_STACKED,
        XL_CHART_TYPE.COLUMN_STACKED_100,
        XL_CHART_TYPE.BAR_STACKED,
        XL_CHART_TYPE.BAR_STACKED_100,
    ]


def is_100_percent_stacked_chart(chart_type):
    return chart_type in [
        XL_CHART_TYPE.COLUMN_STACKED_100,
        XL_CHART_TYPE.BAR_STACKED_100,
    ]


def should_hide_label_dynamic(
    value,
    idx,
    processed_series,
    chart_type,
    series_name=None,
    min_percent_value=0.04,
    min_share_of_total=0.04,
):
    if value is None:
        return True

    # Саму серию total не скрываем этой логикой
    if series_name is not None and str(series_name).strip().lower() == "total":
        return False

    # Для stacked-графиков, включая 100% stacked, считаем долю значения
    # от total конкретной категории. Важно: сами данные НЕ нормализуем,
    # чтобы в embedded Excel под графиком оставались абсолютные значения.
    if is_stacked_chart(chart_type):
        category_total = 0

        for other_series_name, values in processed_series.items():
            if _is_helper_chart_series(other_series_name):
                continue

            if idx < len(values) and values[idx] is not None:
                category_total += values[idx]

        if category_total == 0:
            return True

        threshold = min_percent_value if is_100_percent_stacked_chart(chart_type) else min_share_of_total
        return value / category_total < threshold

    return False


def get_numeric_values(values):
    result = []

    for value in values:
        if value is None:
            result.append(None)

        elif isinstance(value, str) and "%" in value:
            try:
                result.append(float(value.replace("%", "").strip()) / 100)
            except ValueError:
                result.append(None)

        elif isinstance(value, (int, float)):
            result.append(float(value))

        else:
            result.append(None)

    return result


def get_chart_axis_raw_max(series_data, chart_type):
    processed = {
        series_name: get_numeric_values(values)
        for series_name, values in series_data.items()
    }

    stacked_chart_types = [
        XL_CHART_TYPE.COLUMN_STACKED,
        XL_CHART_TYPE.COLUMN_STACKED_100,
        XL_CHART_TYPE.BAR_STACKED,
        XL_CHART_TYPE.BAR_STACKED_100,
    ]

    # Если есть total — берём его как реальную высоту stacked-столбца
    for series_name, values in processed.items():
        if str(series_name).strip().lower() == "total":
            clean_values = [v for v in values if v is not None]
            return max(clean_values) if clean_values else 0

    # Если total нет, для stacked считаем сумму по каждой категории
    if chart_type in stacked_chart_types:
        max_len = max((len(values) for values in processed.values()), default=0)
        category_totals = []

        for idx in range(max_len):
            category_total = 0

            for series_name, values in processed.items():
                if idx < len(values) and values[idx] is not None:
                    category_total += values[idx]

            category_totals.append(category_total)

        return max(category_totals) if category_totals else 0

    # Для обычных графиков берём максимум среди всех значений
    all_values = [
        value
        for values in processed.values()
        for value in values
        if value is not None
    ]

    return max(all_values) if all_values else 0

def get_visible_chart_axis_max(processed_series, chart_type, exclude_series=("total",)):
    exclude_series = {str(x).strip().lower() for x in exclude_series}

    visible_series = {
        series_name: values
        for series_name, values in processed_series.items()
        if str(series_name).strip().lower() not in exclude_series
    }

    if not visible_series:
        return 0

    if is_stacked_chart(chart_type):
        max_len = max((len(values) for values in visible_series.values()), default=0)
        category_totals = []

        for idx in range(max_len):
            category_total = 0

            for values in visible_series.values():
                if idx < len(values) and values[idx] is not None:
                    category_total += values[idx]

            category_totals.append(category_total)

        return max(category_totals) if category_totals else 0

    all_values = [
        value
        for values in visible_series.values()
        for value in values
        if value is not None
    ]

    return max(all_values) if all_values else 0


def get_label_text_color(chart_type, fill_color):
    """
    Для stacked-графиков подписи внутри сегментов — цвет по контрасту.
    Для обычных column/bar графиков подписи обычно снаружи — всегда черные.
    """

    stacked_chart_types = [
        XL_CHART_TYPE.COLUMN_STACKED,
        XL_CHART_TYPE.COLUMN_STACKED_100,
        XL_CHART_TYPE.BAR_STACKED,
        XL_CHART_TYPE.BAR_STACKED_100,
    ]

    if chart_type in stacked_chart_types:
        return (0, 0, 0) if is_bright_color(fill_color) else (255, 255, 255)

    return (0, 0, 0)


def round_up_to_nice_number(value):
    """
    Округляет значение вверх до ближайшего "красивого" числа, максимально близкого к значению
    """
    if value <= 0:
        return 1.0
    
    # Находим порядок числа
    log_val = math.log10(value)
    exp = math.floor(log_val)
    power_of_10 = 10 ** exp
    
    # Нормализуем значение
    normalized = value / power_of_10

    
    # Попробуем использовать стандартный подход, но с меньшими шагами
    if exp >= 3:  # 1000 и больше

        
        rounding_power = exp - 1
        if rounding_power < 0:
            rounding_power = 0
            
        rounding_base = 10 ** rounding_power
        
        # Округляем вверх до ближайшего кратного rounding_base
        rounded_value = math.ceil(value / rounding_base) * rounding_base
        
        # Теперь сделаем красивое округление
        # Проверим, насколько близко к следующему порядку
        next_order = 10 ** exp
        next_nice = 2 * next_order if normalized < 2 else (5 * next_order if normalized < 5 else 10 * next_order)
        
        # Но для простоты будем использовать стандартную логику
        # Округляем вверх до ближайшего красивого числа на нужном уровне
        
        # Попробуем другой подход: округляем до ближайшего красивого числа, но с учетом 2% запаса
        # Для 4233 ищем: 4250, 4300, 4400, 4500, 5000...
        
        # Просто округляем до ближайшего красивого уровня
        # 4233 -> 4300 (округление до следующей сотни)
        
        # Попробуем так: разделим на 10^(exp-1), округлим вверх, умножим обратно
        divisor = 10 ** (exp - 1)  # для 4233 это 100
        quotient = value / divisor  # 42.33
        rounded_quotient = math.ceil(quotient)  # 43
        result = rounded_quotient * divisor  # 43 * 100 = 4300
        
        return result
    else:
        # Для меньших значений используем стандартную логику
        nice_options = [1, 2, 5, 10]
        for nice in nice_options:
            if nice >= normalized:
                return nice * power_of_10
        
        return 10 * power_of_10
    
    
def has_meaningful_values(values):
    return any(v not in ('', None, 0) for v in values)


def _coerce_chart_value(value):
    if value is None:
        return None

    if isinstance(value, str):
        stripped = value.strip().replace(" ", "").replace(",", ".")
        if not stripped:
            return None
        try:
            if stripped.endswith("%"):
                return float(stripped[:-1]) / 100
            return float(stripped)
        except ValueError:
            return None

    if isinstance(value, Number):
        try:
            if math.isnan(value):
                return None
        except TypeError:
            pass
        return float(value)

    return None


def _normalize_series_name(series_name):
    """
    Нормализует имя серии графика.

    Важно: после reset_index / MultiIndex / rename имя серии иногда приходит не как
    чистая строка "total", а как tuple-подобная строка: "('total', '')",
    "total  " или "TOTAL". Поэтому чистим кавычки, скобки и лишние пробелы.
    """
    if isinstance(series_name, (tuple, list)):
        parts = [str(x) for x in series_name if str(x).strip() not in {"", "nan", "None"}]
        text = " ".join(parts)
    else:
        text = str(series_name)

    text = text.strip().lower().replace("ё", "е")
    text = re.sub(r"[\[\]\(\)\{\}'\"`]+", " ", text)
    text = re.sub(r"[,;|]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _series_name_tokens(series_name):
    normalized = _normalize_series_name(series_name)
    if normalized == "":
        return set()
    return set(normalized.split())


def _is_total_series(series_name):
    normalized = _normalize_series_name(series_name)
    tokens = _series_name_tokens(series_name)
    return (
        normalized in {"total", "итого", "всего"}
        or bool(tokens & {"total", "итого", "всего"})
    )


def _is_blank_separator_series(series_name):
    normalized = _normalize_series_name(series_name)
    return normalized in {"", "separator", "разделитель"}


def _is_avg_duration_series(series_name):
    normalized = _normalize_series_name(series_name)
    return normalized in {
        "ср.хроно",
        "ср. хроно",
        "ср хроно",
        "ср.хрон",
        "ср хроно.",
        "avg",
        "avg_duration",
        "average",
    }



DURATION_SERIES_DEFAULT_COLORS = {
    "5": (181, 230, 162),       # #B5E6A2
    "10": (153, 204, 0),        # #99CC00
    "15": (191, 143, 0),        # #BF8F00
    "20": (248, 203, 173),      # #F8CBAD
    "25": (155, 194, 230),      # #9BC2E6
    "30": (47, 85, 151),        # #2F5597
    "50": (47, 85, 151),        # В шаблоне 50 используется как длинные ролики
    "более 30": (255, 153, 204),# #FF99CC
}


def _normalize_duration_series_key(series_name):
    """Нормализует названия серий длительности: '5"' -> '5', 5.0 -> '5'."""
    text = str(series_name).strip().lower().replace("ё", "е")
    text = text.replace('"', '').replace("”", '').replace("“", '').strip()

    try:
        number = float(text.replace(',', '.'))
        if number.is_integer():
            return str(int(number))
        return str(number)
    except Exception:
        return text


def get_default_duration_series_color(series_name):
    """Возвращает дефолтный цвет для duration-серий 5/10/15/20/25/50."""
    return DURATION_SERIES_DEFAULT_COLORS.get(
        _normalize_duration_series_key(series_name)
    )


DURATION_CHART_SERIES_ORDER = ["5", "10", "15", "20", "25", "", "50", "Ср.Хроно"]


def _add_values_elementwise(left, right):
    """Складывает два списка значений поэлементно, сохраняя длину результата."""
    max_len = max(len(left), len(right))
    result = []

    for idx in range(max_len):
        l_val = left[idx] if idx < len(left) else 0
        r_val = right[idx] if idx < len(right) else 0

        l_val = 0 if l_val is None else l_val
        r_val = 0 if r_val is None else r_val

        result.append(l_val + r_val)

    return result


def _duration_series_bucket(series_name):
    """
    Возвращает каноническую серию duration-графика.

    4, 5 -> "5"
    10 -> "10"
    15 -> "15"
    20 -> "20"
    25 -> "25"
    всё больше 25 -> "50"
    Ср.Хроно / avg_duration -> "Ср.Хроно"
    пустая серия -> ""
    """
    if _is_blank_separator_series(series_name):
        return ""

    if _is_avg_duration_series(series_name):
        return "Ср.Хроно"

    key = _normalize_duration_series_key(series_name)

    try:
        value = float(str(key).replace(',', '.'))
    except Exception:
        return None

    if value <= 5:
        return "5"
    if value <= 10:
        return "10"
    if value <= 15:
        return "15"
    if value <= 20:
        return "20"
    if value <= 25:
        return "25"

    # В шаблоне эта серия используется как bucket для длинных роликов.
    return "50"


def prepare_duration_chart_series(processed_series, categories=None):
    """
    Приводит данные duration-графика к точному порядку серий шаблона:
    5, 10, 15, 20, 25, '', 50, Ср.Хроно.

    Это критично для слайдов 13-14: в шаблоне Ср.Хроно находится во втором
    plot, а первые 7 серий — в 100% stacked plot. Если передать другой набор
    или другой порядок серий, PowerPoint начинает красить весь stacked-график
    как Ср.Хроно.
    """
    if categories is not None:
        expected_len = len(categories)
    else:
        expected_len = max((len(v) for v in processed_series.values()), default=0)

    prepared = {
        name: [0] * expected_len
        for name in DURATION_CHART_SERIES_ORDER
    }

    for series_name, values in processed_series.items():
        bucket = _duration_series_bucket(series_name)

        if bucket is None:
            # Для duration-графиков неизвестные серии не передаём в PowerPoint,
            # иначе нарушится структура комбинированного графика из шаблона.
            continue

        clean_values = [0 if value is None else value for value in values]

        if bucket == "Ср.Хроно":
            # Средний хронометраж не суммируем с другими сериями.
            if len(clean_values) < expected_len:
                clean_values = clean_values + [0] * (expected_len - len(clean_values))
            prepared[bucket] = clean_values[:expected_len]
        elif bucket == "":
            # Разделитель должен остаться нулевым.
            prepared[bucket] = [0] * expected_len
        else:
            prepared[bucket] = _add_values_elementwise(prepared[bucket], clean_values)[:expected_len]

    return prepared


def _chart_has_avg_duration_series(chart):
    try:
        return any(_is_avg_duration_series(series.name) for series in chart.series)
    except Exception:
        return False


def _is_helper_chart_series(series_name):
    normalized = _normalize_series_name(series_name)
    return (
        _is_total_series(series_name)
        or _is_blank_separator_series(series_name)
        or _is_avg_duration_series(series_name)
        or normalized in {"", "separator", "разделитель"}
    )




def _set_or_add_chart_bool(parent, tag_name, value):
    """
    Создаёт или обновляет chart XML boolean node, например c:showVal.
    """
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'

    elem = parent.find(f'{{{ns_c}}}{tag_name}')

    if elem is None:
        elem = etree.SubElement(parent, f'{{{ns_c}}}{tag_name}')

    elem.set('val', '1' if value else '0')
    return elem


# Порядок дочерних элементов по схеме OpenXML (DrawingML / Chart).
# Нужен, чтобы вставлять новые узлы НА ПРАВИЛЬНОЕ МЕСТО, а не в конец родителя:
# PowerPoint строго проверяет порядок и молча игнорирует элементы не на своём месте
# (из-за этого total оставался цветным и висел в легенде).
_SPPR_CHILD_ORDER = (
    "xfrm", "custGeom", "prstGeom",
    "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
    "ln", "effectLst", "effectDag", "scene3d", "sp3d", "extLst",
)
_LN_CHILD_ORDER = (
    "noFill", "solidFill", "gradFill", "pattFill",
    "prstDash", "custDash", "round", "bevel", "miter",
    "headEnd", "tailEnd", "extLst",
)
_LEGEND_CHILD_ORDER = (
    "legendPos", "legendEntry", "layout", "overlay", "spPr", "txPr", "extLst",
)
# Общий префикс порядка детей c:ser (одинаков для bar/line/area/pie серий).
_SER_CHILD_ORDER = (
    "idx", "order", "tx", "spPr", "invertIfNegative", "pictureOptions",
    "dPt", "dLbls", "trendline", "errBars",
    "cat", "val", "xVal", "yVal", "bubbleSize", "smooth", "shape", "extLst",
)


def _insert_element_in_order(parent, new_elem, order):
    """
    Вставляет new_elem в parent так, чтобы дети шли в порядке схемы `order`.
    Если место не найдено — добавляет в конец.
    """
    new_local = new_elem.tag.split('}', 1)[-1]
    try:
        new_rank = order.index(new_local)
    except ValueError:
        parent.append(new_elem)
        return new_elem

    for child in parent:
        child_local = child.tag.split('}', 1)[-1]
        if child_local in order and order.index(child_local) > new_rank:
            child.addprevious(new_elem)
            return new_elem

    parent.append(new_elem)
    return new_elem


# Группировки, при которых PowerPoint отображает легенду в ОБРАТНОМ порядке
# (последняя серия — сверху), чтобы порядок легенды совпадал с порядком стопки.
_STACKED_GROUPINGS = {'stacked', 'percentStacked'}
_LEGEND_REVERSING_PLOTS = {'barChart', 'bar3DChart', 'areaChart', 'area3DChart'}


def _legend_is_reversed_for_series(series, ns_c):
    """True, если легенда графика, которому принадлежит серия, отображается
    в обратном порядке (stacked / percentStacked bar|area). В этом случае
    PowerPoint адресует c:legendEntry по ВИЗУАЛЬНОЙ позиции сверху вниз, а не
    по c:idx серии."""
    plot = series._element.getparent()
    if plot is None:
        return False
    local = plot.tag.split('}')[-1]
    if local not in _LEGEND_REVERSING_PLOTS:
        return False
    grouping = plot.find(f'{{{ns_c}}}grouping')
    return grouping is not None and grouping.get('val') in _STACKED_GROUPINGS


def _series_plot_order(series, ns_c):
    """Значение c:order серии (с запасным вариантом c:idx)."""
    e = series._element.find(f'{{{ns_c}}}order')
    if e is not None and e.get('val') is not None:
        return int(e.get('val'))
    e = series._element.find(f'{{{ns_c}}}idx')
    return int(e.get('val')) if (e is not None and e.get('val') is not None) else 0


def reorder_stacked_brands_to_match_forward_legend(chart):
    """У stacked-графиков PowerPoint отображает легенду в обратном порядке
    (последняя серия — сверху). Чтобы нижние графики читались так же, как
    верхний (Магнит сверху, далее по убыванию), разворачиваем порядок брендов
    через c:order. Служебные серии (total и т.п.) оставляем НАВЕРХУ стопки,
    чтобы подпись суммы осталась над столбцом. c:idx (и, значит, цвета и
    привязка данных) не меняются — правим только c:order.
    """
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    series_list = list(chart.series)
    if not series_list:
        return
    if not _legend_is_reversed_for_series(series_list[0], ns_c):
        return

    def _idx(s):
        e = s._element.find(f'{{{ns_c}}}idx')
        return int(e.get('val')) if (e is not None and e.get('val') is not None) else 0

    def _set_order(s, val):
        e = s._element.find(f'{{{ns_c}}}order')
        if e is None:
            idx_el = s._element.find(f'{{{ns_c}}}idx')
            e = etree.Element(f'{{{ns_c}}}order')
            if idx_el is not None:
                idx_el.addnext(e)
            else:
                s._element.insert(0, e)
        e.set('val', str(val))

    helpers = []
    brands = []
    for s in series_list:
        if _is_total_series(s.name) or _is_blank_separator_series(s.name):
            helpers.append(s)
        else:
            brands.append(s)

    # Разворачиваем только те графики, где есть служебная серия (композиция
    # брендов с total). Прочие stacked-графики не трогаем.
    if not helpers or not brands:
        return

    brands.sort(key=_idx)   # по возрастанию c:idx: Магнит первый
    helpers.sort(key=_idx)
    n_brands = len(brands)

    # Бренды занимают нижние слоты порядка 0..n_brands-1; Магнит получает
    # самый высокий из них (окажется сверху видимой легенды). Служебные серии
    # получают верхние слоты и остаются на вершине стопки.
    for rank, s in enumerate(brands):
        _set_order(s, n_brands - 1 - rank)
    for j, s in enumerate(helpers):
        _set_order(s, n_brands + j)


def hide_legend_entries_by_series_names(chart, series_names_to_hide):
    """
    Скрывает отдельные элементы легенды, не удаляя серии из графика.

    Работает не только для точного имени "total", но и для вариантов,
    которые появляются после MultiIndex/reset_index: "('total', '')", "TOTAL" и т.п.
    """
    if not getattr(chart, "has_legend", False):
        return

    normalized_to_hide = {_normalize_series_name(x) for x in series_names_to_hide}
    if not normalized_to_hide:
        return

    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'

    try:
        legend = chart.legend._element
    except Exception:
        return

    # удаляем старые legendEntry, которые могли остаться после повторных запусков
    for old_entry in list(legend.findall(f'{{{ns_c}}}legendEntry')):
        legend.remove(old_entry)

    # Для развёрнутой (stacked) легенды индекс записи считается по визуальной
    # позиции сверху вниз, поэтому нужны общее число серий и их порядок.
    _series_list = list(chart.series)
    _n_entries = len(_series_list)
    _orders_sorted = sorted(_series_plot_order(s, ns_c) for s in _series_list)

    for series in chart.series:
        normalized = _normalize_series_name(series.name)

        should_hide = (
            normalized in normalized_to_hide
            or _is_total_series(series.name)
            or _is_blank_separator_series(series.name)
        )

        if not should_hide:
            continue

        # ВАЖНО: c:legendEntry сопоставляется с серией по её РЕАЛЬНОМУ c:idx
        # из XML, а НЕ по позиции в enumerate(chart.series). В комбинированных
        # (combo) графиках служебная серия total лежит во втором plot и имеет
        # c:idx больше своей позиции. Если брать позицию, delete попадает не в ту
        # запись: total остаётся в легенде, а вместо него скрывается соседний бренд
        # (например «Пятерочка»). Поэтому читаем c:idx прямо из элемента серии.
        series_idx_elem = series._element.find(f'{{{ns_c}}}idx')
        if series_idx_elem is None or series_idx_elem.get('val') is None:
            continue

        # PowerPoint у stacked-графиков разворачивает легенду: последняя серия
        # (например total) показывается СВЕРХУ. При этом c:legendEntry/c:idx
        # адресует запись по её визуальной позиции сверху вниз, а не по c:idx.
        # Поэтому для таких графиков берём (N-1) - позиция_в_прямом_порядке.
        # Для обычных (clustered и т.п.) графиков используем реальный c:idx.
        if _legend_is_reversed_for_series(series, ns_c):
            _forward_rank = _orders_sorted.index(_series_plot_order(series, ns_c))
            series_idx_val = str(_n_entries - 1 - _forward_rank)
        else:
            series_idx_val = series_idx_elem.get('val')

        legend_entry = etree.Element(f'{{{ns_c}}}legendEntry')
        idx_elem = etree.SubElement(legend_entry, f'{{{ns_c}}}idx')
        idx_elem.set('val', series_idx_val)
        delete_elem = etree.SubElement(legend_entry, f'{{{ns_c}}}delete')
        delete_elem.set('val', '1')
        # ВАЖНО: legendEntry должен идти сразу после legendPos, до layout/overlay/spPr/txPr.
        _insert_element_in_order(legend, legend_entry, _LEGEND_CHILD_ORDER)


def hide_series_data_labels(series):
    """Полностью отключает подписи данных у служебной серии."""
    reset_data_labels_xml(series)
    series.has_data_labels = False

    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    ser = series._element
    dLbls = ser.find(f'{{{ns_c}}}dLbls')
    if dLbls is not None:
        for old_dlbl in list(dLbls.findall(f'{{{ns_c}}}dLbl')):
            dLbls.remove(old_dlbl)
        _set_or_add_chart_bool(dLbls, 'showVal', False)
        _set_or_add_chart_bool(dLbls, 'showPercent', False)
        _set_or_add_chart_bool(dLbls, 'showCatName', False)
        _set_or_add_chart_bool(dLbls, 'showSerName', False)
        _set_or_add_chart_bool(dLbls, 'showLegendKey', False)


def _remove_children_by_local_names(element, local_names):
    """Удаляет XML-детей по localname, не завися от namespace-prefix."""
    if element is None:
        return

    for child in list(element):
        local_name = child.tag.split('}', 1)[-1]
        if local_name in local_names:
            element.remove(child)


def remove_broken_ref_series(chart):
    """
    Удаляет из графика «битые» серии шаблона, у которых ссылка указывает на
    #REF! (например <c:f>Лист1!#REF!</c:f> или кэш-имя '#REF!').

    Такие серии остаются в шаблоне после ручного редактирования в PowerPoint
    и при replace_data приводят к повреждению файла: PowerPoint просит
    восстановление и заменяет слайд пустым. Встречается на слайдах 24 и 30.

    Серию не удаляем, если она единственная (чтобы не оставить график без данных).
    """
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    chart_space = chart._chartSpace

    all_sers = list(chart_space.iter(f'{{{ns_c}}}ser'))
    if len(all_sers) <= 1:
        return 0

    removed = 0
    for ser in all_sers:
        is_broken = False

        # 1) Любая формула-ссылка с #REF!
        for f in ser.iter(f'{{{ns_c}}}f'):
            if f.text and '#REF!' in f.text:
                is_broken = True
                break

        # 2) Кэшированное имя серии = '#REF!'
        if not is_broken:
            tx_v = ser.find(f'{{{ns_c}}}tx//{{{ns_c}}}v')
            if tx_v is not None and tx_v.text and '#REF!' in tx_v.text:
                is_broken = True

        if is_broken:
            parent = ser.getparent()
            if parent is not None:
                parent.remove(ser)
                removed += 1

    return removed


def make_series_invisible(series):
    """
    Делает серию визуально невидимой, но не удаляет её из chart data.

    Для служебных серий вроде total и Ср.Хроно это важно:
    - данные и подписи могут остаться;
    - сам столбец/линия не должен быть виден.
    """
    # python-pptx уровень
    try:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor(255, 255, 255)
    except Exception:
        pass

    try:
        series.format.line.width = Pt(0)
        series.format.line.color.rgb = RGBColor(255, 255, 255)
    except Exception:
        pass

    # XML уровень: a:noFill для заливки и линии.
    # Это надёжнее убирает красные рамки у служебных серий в некоторых шаблонах.
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    try:
        ser = series._element
        spPr = ser.find(f'{{{ns_c}}}spPr')
        if spPr is None:
            spPr = etree.Element(f'{{{ns_c}}}spPr')
            _insert_element_in_order(ser, spPr, _SER_CHILD_ORDER)

        _remove_children_by_local_names(
            spPr,
            {'noFill', 'solidFill', 'gradFill', 'pattFill', 'blipFill'}
        )
        # noFill (заливка) должен идти ДО ln/effectLst, иначе PowerPoint его игнорирует
        # и серия рисуется цветом по умолчанию.
        no_fill = etree.Element(f'{{{ns_a}}}noFill')
        _insert_element_in_order(spPr, no_fill, _SPPR_CHILD_ORDER)

        ln = spPr.find(f'{{{ns_a}}}ln')
        if ln is None:
            ln = etree.Element(f'{{{ns_a}}}ln')
            _insert_element_in_order(spPr, ln, _SPPR_CHILD_ORDER)

        _remove_children_by_local_names(
            ln,
            {'noFill', 'solidFill', 'gradFill', 'pattFill', 'blipFill'}
        )
        ln_no_fill = etree.Element(f'{{{ns_a}}}noFill')
        _insert_element_in_order(ln, ln_no_fill, _LN_CHILD_ORDER)
        ln.set('w', '0')
    except Exception:
        pass


def set_data_labels_as_values(data_labels, number_format="#,##0"):
    """
    Обычные графики: показываем значения из chart data.
    """
    dLbls = data_labels._element

    _set_or_add_chart_bool(dLbls, 'showVal', True)
    _set_or_add_chart_bool(dLbls, 'showPercent', False)
    _set_or_add_chart_bool(dLbls, 'showCatName', False)
    _set_or_add_chart_bool(dLbls, 'showSerName', False)
    _set_or_add_chart_bool(dLbls, 'showLegendKey', False)

    try:
        data_labels.number_format = number_format
    except Exception:
        pass


def set_data_labels_as_custom_text(data_labels):
    """
    Для 100% stacked: выключаем автоматический вывод значений.
    Проценты будут записаны вручную в индивидуальные dLbl/tx/rich.

    Это нужно, чтобы:
    - в embedded Excel оставались абсолютные значения;
    - на слайде отображались проценты.
    """
    dLbls = data_labels._element

    _set_or_add_chart_bool(dLbls, 'showVal', False)
    _set_or_add_chart_bool(dLbls, 'showPercent', False)
    _set_or_add_chart_bool(dLbls, 'showCatName', False)
    _set_or_add_chart_bool(dLbls, 'showSerName', False)
    _set_or_add_chart_bool(dLbls, 'showLegendKey', False)


def get_stacked_category_totals(processed_series):
    """
    Считает сумму видимых серий по каждой категории.
    Служебные серии total/итого/... исключаются.
    """
    visible_series = {
        series_name: values
        for series_name, values in processed_series.items()
        if not _is_helper_chart_series(series_name)
    }

    max_len = max((len(values) for values in visible_series.values()), default=0)
    category_totals = []

    for idx in range(max_len):
        total = 0

        for values in visible_series.values():
            if idx < len(values) and values[idx] is not None:
                total += values[idx]

        category_totals.append(total)

    return category_totals


def add_custom_percent_labels_to_series(
    series,
    values,
    category_totals,
    min_percent_to_show=0.02,
):
    """
    Создаёт индивидуальные подписи точек для 100% stacked-графика.

    Данные в графике остаются абсолютными, но подписи на слайде показываются
    как проценты от суммы категории.
    """
    ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    ser = series._element
    dLbls = ser.find(f'{{{ns_c}}}dLbls')

    if dLbls is None:
        dLbls = etree.SubElement(ser, f'{{{ns_c}}}dLbls')

    # Удаляем старые индивидуальные подписи, чтобы при повторном запуске
    # не плодились дубликаты.
    for old_dlbl in list(dLbls.findall(f'{{{ns_c}}}dLbl')):
        dLbls.remove(old_dlbl)

    _set_or_add_chart_bool(dLbls, 'showVal', False)
    _set_or_add_chart_bool(dLbls, 'showPercent', False)
    _set_or_add_chart_bool(dLbls, 'showCatName', False)
    _set_or_add_chart_bool(dLbls, 'showSerName', False)
    _set_or_add_chart_bool(dLbls, 'showLegendKey', False)

    for idx, value in enumerate(values):
        if value is None or idx >= len(category_totals):
            continue

        total = category_totals[idx]

        if total in [None, 0]:
            continue

        share = value / total

        if share < min_percent_to_show:
            continue

        label_text = f"{share:.0%}"

        dLbl = etree.SubElement(dLbls, f'{{{ns_c}}}dLbl')

        idx_elem = etree.SubElement(dLbl, f'{{{ns_c}}}idx')
        idx_elem.set('val', str(idx))

        tx = etree.SubElement(dLbl, f'{{{ns_c}}}tx')
        rich = etree.SubElement(tx, f'{{{ns_c}}}rich')

        body_pr = etree.SubElement(rich, f'{{{ns_a}}}bodyPr')
        body_pr.set('wrap', 'none')
        body_pr.set('lIns', '0')
        body_pr.set('rIns', '0')
        body_pr.set('tIns', '0')
        body_pr.set('bIns', '0')

        etree.SubElement(rich, f'{{{ns_a}}}lstStyle')

        p = etree.SubElement(rich, f'{{{ns_a}}}p')
        r = etree.SubElement(p, f'{{{ns_a}}}r')

        rPr = etree.SubElement(r, f'{{{ns_a}}}rPr')
        rPr.set('lang', 'ru-RU')

        t = etree.SubElement(r, f'{{{ns_a}}}t')
        t.text = label_text

        etree.SubElement(p, f'{{{ns_a}}}endParaRPr')

        _set_or_add_chart_bool(dLbl, 'showVal', False)
        _set_or_add_chart_bool(dLbl, 'showPercent', False)
        _set_or_add_chart_bool(dLbl, 'showCatName', False)
        _set_or_add_chart_bool(dLbl, 'showSerName', False)
        _set_or_add_chart_bool(dLbl, 'showLegendKey', False)


def find_shape_by_name(slide, shape_name):
    """
    Ищет shape по имени на слайде, включая grouped shapes.
    Имя shape задаётся в PowerPoint через: Главная -> Выделить -> Область выделения.
    """
    for shape in iter_all_shapes(slide.shapes):
        if shape.name == shape_name:
            return shape
    return None


def update_chart_data(
    presentation,
    slide,
    chart_id=None,
    chart_type=None,
    categories=None,
    series_data=None,
    color_dict=None,
    y_for_percentage=None,
    size_label=10,
    custom_y_axis_max=None,
    y_axis_margin=None,
    exclude_total_from_axis=True,
    y_axis_margin_rules=None,
    shape_name=None,
    shape_id=None,
    stacked_value_labels_as_percent=False,
):
    slide_num = slide
    slide = presentation.slides[slide - 1]
    count = 0

    raw_processed_series = {
        series_name: [_coerce_chart_value(value) for value in values]
        for series_name, values in series_data.items()
    }

    # ВАЖНО: для COLUMN_STACKED_100 / BAR_STACKED_100 НЕ нормализуем данные вручную.
    # PowerPoint сам отображает 100%-stacked как доли, даже если в embedded Excel
    # лежат абсолютные значения. Если нормализовать здесь, то под капотом
    # в графике окажутся доли 0..1 вместо исходных абсолютов.
    processed_series = raw_processed_series

    is_100_percent_chart = is_100_percent_stacked_chart(chart_type)

    # Duration-графики с Ср.Хроно — комбинированные графики из шаблона.
    # Им нужен точный набор и порядок серий, иначе Ср.Хроно попадает в stacked plot
    # и начинает красить весь график своим цветом.
    input_has_avg_duration_series = any(
        _is_avg_duration_series(series_name)
        for series_name in processed_series.keys()
    )

    if input_has_avg_duration_series:
        processed_series = prepare_duration_chart_series(
            processed_series,
            categories=categories,
        )

    if is_100_percent_chart:
        stacked100_category_totals = get_stacked_category_totals(processed_series)
    else:
        stacked100_category_totals = None

    # Для обычных stacked-графиков, где внутри столбца нужны проценты,
    # а служебная серия total сверху должна остаться абсолютным числом.
    if stacked_value_labels_as_percent:
        stacked_category_totals = get_stacked_category_totals(processed_series)
    else:
        stacked_category_totals = None

    # Duration-графики распознаём по служебной серии Ср.Хроно.
    # Для них автоматически применяем цвета по длительностям 5/10/15/20/25/50.
    has_avg_duration_series = any(
        _is_avg_duration_series(series_name)
        for series_name in processed_series.keys()
    )

    # Определяем тип данных для оси, но не для замены исходных данных.
    is_percentage = is_100_percent_chart or (
        bool(processed_series) and all(
            has_meaningful_values(values) and all(
                isinstance(v, float) and 0 < v <= 1
                for v in values
                if v not in ('', None, 0)
            )
            for values in processed_series.values()
        )
    )

    # Формат оси и формат подписей данных разделяем.
    # Для 100%-stacked ось процентная, но сами значения в chart data остаются абсолютными.
    if is_100_percent_chart:
        axis_number_format = "0%"
        # В embedded Excel лежат абсолютные значения, но на слайде
        # подписи для 100%-stacked будут кастомными процентами.
        label_number_format = "0%"
    elif is_percentage:
        axis_number_format = "0%"
        label_number_format = "0%"
    else:
        axis_number_format = "#,##0"
        label_number_format = "#,##0"

    # Оставляем старое имя для кусков кода, где нужен формат подписей.
    number_format = label_number_format


    is_pie_or_doughnut = chart_type in [XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT]

    for shape in iter_all_shapes(slide.shapes):
        if shape.has_chart:
            count += 1

            is_target_chart = (
                (shape_name is not None and shape.name == shape_name)
                or
                (shape_id is not None and getattr(shape, "shape_id", None) == shape_id)
                or
                (shape_name is None and shape_id is None and chart_id is not None and count == chart_id)
            )

            if is_target_chart:
                chart = shape.chart
                if chart_type is not None and chart.chart_type != chart_type:
                    print(
                        f"Chart type does not match on slide {slide_num}: "
                        f"shape_name={shape.name!r}, chart_id={chart_id}, "
                        f"expected={chart_type}, actual={chart.chart_type}"
                    )
                    return
                
                # Удаляем «битые» служебные серии шаблона (#REF!), иначе
                # сломанная ссылка Лист1!#REF! ломает график в PowerPoint
                # («презентация требует восстановления», слайд становится пустым).
                # Актуально для слайдов 24 и 30.
                remove_broken_ref_series(chart)

                # Обновление данных диаграммы
                chart_data = CategoryChartData()
                chart_data.categories = [str(c) if c is not None else "" for c in categories]  # Обработка None в категориях

                for series_name, values in processed_series.items():
                    chart_data.add_series(series_name, values)


                print(shape.name, list(processed_series.keys()))
                print(shape.name, "категорий:", len(categories))
                for k, v in processed_series.items():
                    print("  ", repr(k), len(v))

                if _chart_has_avg_duration_series(chart):
                    update_chart_values_in_place(chart, processed_series, categories)
                    _update_chart_workbook(chart, categories, processed_series)
                    pass
                else: 
                    chart.replace_data(chart_data)

                # У stacked-графиков разворачиваем порядок брендов, чтобы легенда
                # читалась как на верхнем графике (Магнит сверху). total остаётся
                # наверху стопки. Делать это нужно ДО скрытия легенды.
                if not _chart_has_avg_duration_series(chart):
                    reorder_stacked_brands_to_match_forward_legend(chart)

                # Служебные серии нужны для расчётов/подписей, но не должны попадать в легенду.
                hide_legend_entries_by_series_names(chart, ["", "separator", "разделитель", "total", "итого", "всего"])

                # Настройки для диаграмм с осями (не PIE/DOUGHNUT)
                if not is_pie_or_doughnut:
                    try:
                        value_axis = chart.value_axis
                        
                        # Для 100%-stacked ось всегда 0..100%.
                        # custom_y_axis_max для такого графика игнорируем.
                        if is_100_percent_chart:
                            y_axis_max = 1.0

                        elif custom_y_axis_max is not None:
                            y_axis_max = custom_y_axis_max

                        elif is_percentage:
                            if y_for_percentage:
                                y_axis_max = y_for_percentage
                            else:
                                y_axis_max = 1.0

                        else:
                            has_total_series = "total" in [
                                str(x).strip().lower()
                                for x in processed_series.keys()
                            ]

                            if exclude_total_from_axis:
                                max_value = get_visible_chart_axis_max(
                                    processed_series=processed_series,
                                    chart_type=chart_type,
                                    exclude_series=("total",),
                                )
                            else:
                                max_value = get_visible_chart_axis_max(
                                    processed_series=processed_series,
                                    chart_type=chart_type,
                                    exclude_series=(),
                                )

                            # Базовый запас сохраняет прежнюю логику модуля:
                            # если total служебный и исключён из оси — даём больше воздуха сверху.
                            axis_margin = 1.18 if (has_total_series and exclude_total_from_axis) else 1.02

                            # Старый способ через словарь оставлен для обратной совместимости.
                            if y_axis_margin_rules:
                                axis_margin = y_axis_margin_rules.get((slide_num, chart_id), axis_margin)

                            # Новый способ: точечно передать запас прямо в update_chart_data(...).
                            # Например: y_axis_margin=1.25, exclude_total_from_axis=True
                            if y_axis_margin is not None:
                                axis_margin = y_axis_margin

                            value_with_margin = max_value * axis_margin
                            y_axis_max = round_up_to_nice_number(value_with_margin)

                        # 🔹 Устанавливаем максимум
                        value_axis.maximum_scale = y_axis_max
                        
                        # Ось всегда начинается с 0. Для 100%-stacked это 0%.
                        value_axis.minimum_scale = 0.0

                        value_axis.tick_labels.number_format = axis_number_format
                        
                        # Устанавливаем красивый интервал делений
                        if is_percentage:
                            value_axis.major_unit = 0.2
                        else:
                            # Рассчитываем красивый интервал
                            raw_unit = y_axis_max / 10
                            
                            # Находим красивый интервал
                            log_unit = math.log10(raw_unit)
                            exp = math.floor(log_unit)
                            power_of_10 = 10 ** exp
                            
                            normalized_unit = raw_unit / power_of_10
                            
                            # Выбираем красивое число для интервала
                            nice_options = [1, 2, 5, 10]
                            for nice in nice_options:
                                if nice >= normalized_unit:
                                    value_axis.major_unit = nice * power_of_10
                                    break

                        value_axis.minor_tick_mark = XL_TICK_MARK.NONE
                        if value_axis.tick_labels:
                            value_axis.tick_labels.font.size = Pt(size_label)

                        # --- Фикс «странных осей» на комбо-графиках (слайды 24/30) ---
                        # У таких графиков ДВЕ оси значений (c:valAx): «Кол-во товаров»
                        # и «AWW». python-pptx через chart.value_axis настраивает только
                        # одну из них (valAx_lst[1], если осей две), а серия-столбцы может
                        # быть привязана к другой оси, у которой в шаблоне остаётся
                        # «мусорный» масштаб (например 0..80 или -80..120). Из-за этого
                        # столбцы упираются в верх графика, а ось уходит в минус.
                        # Применяем один и тот же масштаб ко ВСЕМ осям значений, чтобы
                        # картинка не зависела от того, к какой оси привязаны столбцы.
                        try:
                            resolved_major_unit = value_axis.major_unit
                        except Exception:
                            resolved_major_unit = None

                        for _valAx in list(chart._chartSpace.valAx_lst):
                            try:
                                extra_axis = ValueAxis(_valAx)
                                extra_axis.maximum_scale = y_axis_max
                                extra_axis.minimum_scale = 0.0
                                extra_axis.tick_labels.number_format = axis_number_format
                                if resolved_major_unit is not None:
                                    extra_axis.major_unit = resolved_major_unit
                                extra_axis.minor_tick_mark = XL_TICK_MARK.NONE
                                if extra_axis.tick_labels:
                                    extra_axis.tick_labels.font.size = Pt(size_label)
                            except Exception:
                                # Битую/нестандартную ось просто пропускаем.
                                pass

                    except AttributeError:
                        pass

                # Специальные настройки для PIE/DOUGHNUT
                if is_pie_or_doughnut:
                    chart.has_legend = False
                    plot = chart.plots[0]
                    plot.has_data_labels = True
                    data_labels = plot.data_labels
                    data_labels.show_percentage = True
                    data_labels.show_category_name = False
                    data_labels.show_value = False
                    data_labels.font.size = Pt(size_label)
                    configure_data_labels_no_wrap(data_labels)

                    
                    # Настройка цветов для каждой категории
                    for idx, point in enumerate(plot.series[0].points):
                        # Получаем название категории (из списка categories)
                        category_name = categories[idx] if idx < len(categories) else f"Категория {idx}"
                        
                        # Получаем цвет из словаря (по имени категории)
                        color = color_dict.get(str(category_name)) or (128, 128, 128)  # Серый по умолчанию
                        
                        # Применяем цвет к точке
                        fill = point.format.fill
                        fill.solid()
                        fill.fore_color.rgb = RGBColor(*color)

                

                MIN_VALUE_TO_SHOW = 1.0  # Например, не показывать значения меньше 1

                for series in chart.series:
                    series_name_normalized = _normalize_series_name(series.name)

                    # Пустая серия-разделитель: нужна в данных как визуальный gap,
                    # но не должна иметь подписи и не должна быть видна в легенде.
                    if _is_blank_separator_series(series.name):
                        make_series_invisible(series)
                        hide_series_data_labels(series)
                        continue

                    if _is_total_series(series.name):
                        reset_data_labels_xml(series)
                        series.has_data_labels = True
                        data_labels = series.data_labels
                        
                        # Устанавливаем формат и стиль меток
                        data_labels.position = XL_LABEL_POSITION.INSIDE_BASE
                        data_labels.font.size = Pt(size_label)
                        data_labels.font.bold = True
                        data_labels.font.color.rgb = RGBColor(127, 127, 177)
                        data_labels.number_format = number_format
                        
                        # Включаем показ значений (но позже отключим для маленьких)
                        data_labels.show_value = True
                        configure_data_labels_no_wrap(data_labels)
                        
                        # Серия total служебная: столбец должен быть невидимым,
                        # но подпись total сверху остаётся видимой.
                        make_series_invisible(series)
                        
                        # Проходим по точкам данных и скрываем маленькие значения
                        for point in series.points:
                            if point.data_label:
                                # Получаем значение точки (если возможно)
                                try:
                                    value = float(point.data_label.text)  # Пытаемся преобразовать текст метки в число
                                except (ValueError, AttributeError):
                                    continue  # Если не число, пропускаем
                                
                                # Если значение меньше порога, скрываем метку
                                if value < MIN_VALUE_TO_SHOW:
                                    point.data_label.text = ""  # Очищаем текст
                                    point.data_label.has_text_frame = False  # Отключаем отображение

                    # Настройка для обычных серий
                    else:
                        reset_data_labels_xml(series)
                        series.has_data_labels = True
                        data_labels = series.data_labels
                        
                        if chart_type not in [XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS]:
                            data_labels.show_value = True
                        else:
                            data_labels.show_value = False
                        
                        # Явно отключаем лишнее
                        data_labels.show_category_name = False
                        data_labels.show_series_name = False
                        data_labels.show_percentage = False
                        
                        data_labels.font.size = Pt(size_label)
                        data_labels.font.bold = False

                        # Ср.Хроно — служебная серия комбинированных duration-графиков.
                        # Данные в Excel оставляем, подписи показываем числом,
                        # но сам столбец/линия делаем невидимым, чтобы окрашивались именно
                        # серии длительностей 5/10/15/20/25/50.
                        if _is_avg_duration_series(series.name):
                            # Ср.Хроно — служебная серия второго plot.
                            # Подписи оставляем числом, но саму серию делаем невидимой,
                            # чтобы цвета применялись только к сериям секунд.
                            set_data_labels_as_values(data_labels, number_format="0.0")
                            configure_data_labels_no_wrap(data_labels)

                            try:
                                data_labels.position = XL_LABEL_POSITION.CENTER
                            except Exception:
                                pass

                            try:
                                data_labels.font.color.rgb = RGBColor(0, 0, 0)
                            except Exception:
                                pass

                            make_series_invisible(series)
                            continue

                        if is_100_percent_chart or stacked_value_labels_as_percent:
                            set_data_labels_as_custom_text(data_labels)
                        else:
                            set_data_labels_as_values(data_labels, number_format=label_number_format)

                        configure_data_labels_no_wrap(data_labels)

                        if is_100_percent_chart:
                            series_values = processed_series.get(str(series.name), [])
                            add_custom_percent_labels_to_series(
                                series=series,
                                values=series_values,
                                category_totals=stacked100_category_totals,
                                min_percent_to_show=0.02,
                            )

                        elif stacked_value_labels_as_percent:
                            series_values = processed_series.get(str(series.name), [])
                            add_custom_percent_labels_to_series(
                                series=series,
                                values=series_values,
                                category_totals=stacked_category_totals,
                                min_percent_to_show=0.02,
                            )

                        fill_color = None
                        series_name_text = str(series.name)

                        if color_dict:
                            if series_name_text in color_dict:
                                fill_color = color_dict[series_name_text]
                            else:
                                # Дополнительные варианты для duration-серий:
                                # 5, 5", 5.0 должны находить один и тот же цвет.
                                normalized_duration_key = _normalize_duration_series_key(series_name_text)
                                for candidate in (
                                    normalized_duration_key,
                                    f'{normalized_duration_key}"',
                                    normalized_duration_key.upper(),
                                ):
                                    if candidate in color_dict:
                                        fill_color = color_dict[candidate]
                                        break

                        if fill_color is None and has_avg_duration_series:
                            fill_color = get_default_duration_series_color(series_name_text)

                        if fill_color:
                            text_color = get_label_text_color(chart_type, fill_color)
                            data_labels.font.color.rgb = RGBColor(*text_color)

                            if chart_type in [XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS]:
                                series.format.line.color.rgb = RGBColor(*fill_color)
                            else:
                                series.format.fill.solid()
                                series.format.fill.fore_color.rgb = RGBColor(*fill_color)
                                series.format.line.width = Pt(0)  # убираем линию для bar/column

                        # Настройки индивидуальных подписей точек.
                        # Для 100%-stacked подписи уже созданы вручную как проценты,
                        # поэтому не включаем show_value и не запускаем старую логику скрытия.
                        if not (is_100_percent_chart or stacked_value_labels_as_percent):
                            for point in series.points:
                                if point.data_label:
                                    point.data_label.show_value = True
                                    point.data_label.show_category_name = False
                                    point.data_label.show_series_name = False
                                    point.data_label.show_percentage = False

                            ns_c = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
                            series_name = str(series.name)

                            if series_name in processed_series:
                                values = processed_series[series_name]

                                for idx, point in enumerate(series.points):
                                    if idx < len(values):
                                        val = values[idx]
                                        
                                        should_hide = should_hide_label_dynamic(
                                            value=val,
                                            idx=idx,
                                            processed_series=processed_series,
                                            chart_type=chart_type,
                                            series_name=series_name,
                                            min_percent_value=0.04,
                                            min_share_of_total=0.04,
                                        )
                                        
                                        if should_hide:
                                            dLbls = series._element.find(f'{{{ns_c}}}dLbls')
                                            if dLbls is not None:
                                                dLbl = etree.SubElement(dLbls, f'{{{ns_c}}}dLbl')
                                                idx_elem = etree.SubElement(dLbl, f'{{{ns_c}}}idx')
                                                idx_elem.set('val', str(idx))
                                                show_val = etree.SubElement(dLbl, f'{{{ns_c}}}showVal')
                                                show_val.set('val', '0')
                                                show_legend = etree.SubElement(dLbl, f'{{{ns_c}}}showLegendKey')
                                                show_legend.set('val', '0')
                                                show_cat = etree.SubElement(dLbl, f'{{{ns_c}}}showCatName')
                                                show_cat.set('val', '0')
                                                show_ser = etree.SubElement(dLbl, f'{{{ns_c}}}showSerName')
                                                show_ser.set('val', '0')

                
                return
                
    if shape_name is not None or shape_id is not None:
        print(f"Chart with shape_name={shape_name!r}, shape_id={shape_id!r} not found on slide {slide_num}.")
    else:
        print(f"Chart #{chart_id} not found on slide {slide_num}.")



def update_chart_data_by_shape_name(
    presentation,
    slide,
    shape_name,
    chart_type,
    categories,
    series_data,
    color_dict=None,
    y_for_percentage=None,
    size_label=10,
    custom_y_axis_max=None,
    y_axis_margin=None,
    exclude_total_from_axis=True,
    shape_id=None,
    stacked_value_labels_as_percent=False,
):
    """
    Новый устойчивый способ обновления графика: по имени shape, а не по chart_id.

    Используй, когда в шаблоне графики названы через Selection Pane.
    Старый update_chart_data(..., chart_id=...) продолжает работать.
    """
    return update_chart_data(
        presentation=presentation,
        slide=slide,
        chart_id=None,
        shape_name=shape_name,
        shape_id=shape_id,
        chart_type=chart_type,
        categories=categories,
        series_data=series_data,
        color_dict=color_dict,
        y_for_percentage=y_for_percentage,
        size_label=size_label,
        custom_y_axis_max=custom_y_axis_max,
        y_axis_margin=y_axis_margin,
        exclude_total_from_axis=exclude_total_from_axis,
        stacked_value_labels_as_percent=stacked_value_labels_as_percent,
    )


def change_data_for_ppt(df, name_column, keep_empty_series=False):
    df = df.copy()

    if name_column not in df.columns:
        raise KeyError(f"В dataframe нет колонки {name_column!r}")

    categories = df[name_column].astype(str).tolist()

    value_df = df.drop(columns=[name_column])

    series_data = {}

    for col in value_df.columns:
        numeric_values = pd.to_numeric(
            value_df[col].replace('', pd.NA),
            errors='coerce'
        )

        # если колонка полностью текстовая — не добавляем её в график
        if numeric_values.notna().sum() == 0:
            if not keep_empty_series:
                continue
            numeric_values = pd.Series([0] * len(value_df), index=value_df.index)

        series_data[str(col)] = [
            None if pd.isna(v) else float(v)
            for v in numeric_values
        ]

    return categories, series_data


def _set_ppt_table_cell_text(cell, text, font_size=8):
    """Пишет текст в ячейку PowerPoint-таблицы и выставляет кегль."""
    cell.text = "" if text is None else str(text)

    try:
        text_frame = cell.text_frame
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
    except Exception:
        pass


# -----------------------------
# Duration tables on slides 13-14
# -----------------------------

def _normalize_duration_row_name(value):
    """
    Нормализует названия строк duration-таблицы:
    'Средний\nХроно', 'СРЕДНИЙ ХРОНО', 'Средний Хроно' -> 'средний хроно'.
    """
    return (
        str(value)
        .replace("\\n", "\n")
        .replace("\n", " ")
        .strip()
        .lower()
        .replace("ё", "е")
    )


def _find_duration_label_column(df):
    """
    Находит колонку с названиями строк duration-таблицы после reset_index().

    Поддерживает варианты:
    - index
    - _period_group
    - первая колонка, где значениями идут Средний\nХроно / Промо / Имидж / Итого.
    """
    target_names = {
        "средний хроно",
        "промо",
        "имидж",
        "итого",
    }

    for col in df.columns:
        if str(col) in {"index", "_period_group", "period_group"}:
            return col

    for col in df.columns:
        try:
            values = {
                _normalize_duration_row_name(v)
                for v in df[col].head(12).tolist()
            }
        except Exception:
            continue

        if values & target_names:
            return col

    return None


def find_first_table_on_slide(slide):
    """Возвращает первую таблицу на слайде, если она не была названа в шаблоне."""
    for shape in iter_all_shapes(slide.shapes):
        if shape.has_table:
            return shape
    return None


# -----------------------------
# Promo-share tables on slide 23
# -----------------------------

def _normalize_promo_label(value):
    """
    Универсальная нормализация подписи (строки/столбца) для promo-таблиц:
    нижний регистр, ё->е, схлопывание пробелов, без хвостовых пробелов/неразрывных
    пробелов. 'Пятёрочка ' / 'ПЯТЕРОЧКА' / 'пятерочка' -> 'пятерочка'.
    """
    text = str(value).replace("\xa0", " ").replace("\n", " ")
    text = text.strip().lower().replace("ё", "е")
    return " ".join(text.split())


# Подписи строк PPT-таблицы -> возможные имена в индексе датафрейма.
# Ключ — то, что написано в шаблоне (нормализованное), значения — алиасы из df.
PROMO_ROW_ALIASES = {
    "пятерочка": {"пятерочка", "x5", "пятёрочка"},
    "магнит": {"магнит"},
    "категория": {
        "категория",
        "promo_share_category",
        "promo share category",
        "promo_share",
        "итого",
        "всего",
    },
}


def _is_month_label(value):
    """Похоже ли значение на подпись месяца вида 'янв.25', 'дек.26'?"""
    text = _normalize_promo_label(value)
    month_prefixes = ("янв", "фев", "мар", "апр", "май", "июн",
                      "июл", "авг", "сен", "окт", "ноя", "дек")
    return text.startswith(month_prefixes) and "." in text


def _promo_value_is_empty(value):
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    text = str(value).strip()
    return text in {"", "nan", "none", "None", "-"}


def _format_promo_value(value, as_percent=True, decimals=0, suffix="%"):
    """
    Форматирует значение promo-доли.
    Доли (0..1) при as_percent=True умножаются на 100.
    Возвращает '' для пустых значений.
    """
    if _promo_value_is_empty(value):
        return ""

    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value).strip()

    if as_percent:
        number *= 100.0

    if number == 0:
        return ""

    text = f"{number:.{decimals}f}".replace(".", ",")
    return f"{text}{suffix}" if suffix else text


def _auto_detect_fraction(df_numeric_values):
    """Если все значения по модулю <= 1.5 — считаем, что это доли (0..1)."""
    finite = [
        abs(float(v))
        for v in df_numeric_values
        if not _promo_value_is_empty(v)
        and isinstance(v, Number)
    ]
    if not finite:
        return False
    return max(finite) <= 1.5


def fill_promo_share_table(
    table,
    df,
    *,
    as_percent="auto",
    decimals=0,
    suffix="%",
    font_size=8,
    empty_text="",
    row_aliases=None,
    verbose=True,
    name="",
):
    """
    Заполняет promo-таблицу на слайде 23 (table__offline_promo_nat / _reg).

    Структура PPT-таблицы (4x16):
        row 0  — шапка: [подпись] + месяцы (янв.25 ... мар.26)
        row 1  — Пятерочка
        row 2  — Магнит
        row 3  — Категория

    Датафрейм (после build_promo_share + rename_month_year):
        index   — бренды + 'promo_share_category' (в любом регистре),
                  либо те же значения в колонке 'index' после reset_index();
        columns — подписи месяцев 'янв.25' ... 'мар.26';
        значения — доли промо (0..1).

    Заливка идёт ПО СОВПАДЕНИЮ ПОДПИСЕЙ:
      - строки сопоставляются через PROMO_ROW_ALIASES (Категория <- promo_share_category);
      - столбцы — по подписи месяца;
    поэтому порядок строк/столбцов в df значения не имеет, а 'Категория'
    подхватывается из 'promo_share_category'.

    as_percent:
        'auto' — определить автоматически (значения <= 1.5 трактуются как доли);
        True   — всегда умножать на 100;
        False  — выводить как есть.
    """
    if row_aliases is None:
        row_aliases = PROMO_ROW_ALIASES

    data = df.copy()

    # Если индекс вынесен в колонку 'index' (как делает общий прогон) — вернём его.
    if "index" in data.columns:
        data = data.set_index("index")

    data.index = data.index.map(str)
    data.columns = data.columns.map(str)

    # Определяем ориентацию: месяцы должны быть в КОЛОНКАХ.
    cols_look_like_months = sum(_is_month_label(c) for c in data.columns)
    rows_look_like_months = sum(_is_month_label(i) for i in data.index)
    if rows_look_like_months > cols_look_like_months:
        data = data.T

    # Карта подписей df -> имя в df (нормализованные).
    df_index_lookup = {_normalize_promo_label(i): i for i in data.index}
    df_col_lookup = {_normalize_promo_label(c): c for c in data.columns}

    # as_percent='auto'
    if as_percent == "auto":
        flat_values = [v for v in data.to_numpy().ravel()]
        as_percent = _auto_detect_fraction(flat_values)

    R, C = len(table.rows), len(table.columns)

    written = 0
    unmatched_rows = []

    for r in range(1, R):  # row 0 — шапка
        ppt_row_label = _normalize_promo_label(table.cell(r, 0).text)

        # Находим имя строки в df через алиасы.
        df_row_name = None
        aliases = row_aliases.get(ppt_row_label, {ppt_row_label})
        for alias in aliases:
            alias_norm = _normalize_promo_label(alias)
            if alias_norm in df_index_lookup:
                df_row_name = df_index_lookup[alias_norm]
                break
        # Fallback: прямое совпадение подписи.
        if df_row_name is None and ppt_row_label in df_index_lookup:
            df_row_name = df_index_lookup[ppt_row_label]

        if df_row_name is None:
            unmatched_rows.append(table.cell(r, 0).text)
            continue

        for c in range(1, C):  # col 0 — подпись строки
            ppt_col_label = _normalize_promo_label(table.cell(0, c).text)
            df_col_name = df_col_lookup.get(ppt_col_label)
            if df_col_name is None:
                continue

            value = data.loc[df_row_name, df_col_name]
            if isinstance(value, pd.Series):
                value = value.dropna()
                value = value.iloc[0] if len(value) else None

            text = empty_text if _promo_value_is_empty(value) else _format_promo_value(
                value, as_percent=as_percent, decimals=decimals, suffix=suffix
            )
            _set_ppt_table_cell_text(table.cell(r, c), text, font_size=font_size)
            written += 1

    if verbose:
        print(f"  [{name}] promo-таблица: записано {written} ячеек, "
              f"as_percent={as_percent}"
              + (f", не сопоставлены строки: {unmatched_rows}" if unmatched_rows else ""))

    return written


def fill_promo_share_table_on_slide(
    presentation,
    slide_num,
    df,
    shape_name=None,
    fallback_to_first_table=True,
    **kwargs,
):
    """
    Удобная обёртка для promo-таблиц на слайде 23.

    Если shape_name задан — ищет таблицу по имени;
    иначе (и при fallback_to_first_table=True) берёт первую таблицу на слайде.
    """
    slide = presentation.slides[slide_num - 1]

    table_shape = None
    if shape_name is not None:
        table_shape = find_shape_by_name(slide, shape_name)
    if table_shape is None and fallback_to_first_table:
        table_shape = find_first_table_on_slide(slide)

    if table_shape is None or not table_shape.has_table:
        raise ValueError(
            f"Не найдена таблица на слайде {slide_num}. shape_name={shape_name!r}"
        )

    fill_promo_share_table(
        table_shape.table,
        df,
        name=kwargs.pop("name", shape_name or f"slide{slide_num}"),
        **kwargs,
    )
    return table_shape


# -----------------------------
# SOS / SOV tables on slides 25 and 31
# -----------------------------

# Русские сокращения месяцев -> номер. Совпадает с rename_month_year ('янв.25').
_RU_MONTH_TO_NUM = {
    "янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
    "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12,
}


def _parse_month_year_label(label):
    """
    'янв.25' -> (month_num=1, year2=25). Возвращает (None, None), если не распознано.
    Понимает разделители '.', '-', ' ' и регистр.
    """
    text = _normalize_promo_label(label)  # нижний регистр, ё->е, без лишних пробелов
    text = text.replace("-", ".").replace(" ", ".")
    parts = [p for p in text.split(".") if p != ""]
    if len(parts) < 2:
        return None, None

    month_num = _RU_MONTH_TO_NUM.get(parts[0][:3])
    year_digits = "".join(ch for ch in parts[-1] if ch.isdigit())
    if month_num is None or year_digits == "":
        return None, None

    year2 = int(year_digits[-2:])
    return month_num, year2


def _detect_year_blocks_from_number_row(table, number_row, label_cols=1):
    """
    Считывает строку с номерами месяцев (например '1'..'12' дважды) и возвращает
    два словаря {month_num: ppt_col}: для левого ({ly}) и правого ({c_year}) блоков.

    Разделение на блоки идёт по «обнулению» последовательности: после первого
    блока 1..12 идёт пустой столбец-разделитель, затем снова 1..12.
    """
    blocks = []
    current = {}
    prev_num = None

    for c in range(label_cols, len(table.columns)):
        raw = table.cell(number_row, c).text.strip()
        if not raw.isdigit():
            # разделитель/пустой столбец завершает текущий блок
            if current:
                blocks.append(current)
                current = {}
            prev_num = None
            continue

        num = int(raw)
        # новый блок, если номер пошёл на убыль (12 -> 1)
        if prev_num is not None and num <= prev_num and current:
            blocks.append(current)
            current = {}

        current[num] = c
        prev_num = num

    if current:
        blocks.append(current)

    left = blocks[0] if len(blocks) >= 1 else {}
    right = blocks[1] if len(blocks) >= 2 else {}
    return left, right


def fill_sos_sov_table(
    table,
    df,
    *,
    number_row=1,
    data_start_row=2,
    label_col=0,
    ly_year=None,
    c_year=None,
    value_format=None,
    as_percent=False,
    decimals=0,
    suffix="",
    empty_text="",
    font_size=7,
    verbose=True,
    name="",
):
    """
    Заполняет SOS/SOV-таблицу со слайдов 25/31.

    Структура PPT-таблицы (6x26 или 5x26):
        row 0      — заголовки секций ('SOS {ly}', 'SOS {c_year}'), объединённые;
        row 1      — номера месяцев: 1..12 для {ly} (левый блок) и 1..12 для {c_year};
        row 2..N   — бренды (ПЯТЕРОЧКА / ПЕРЕКРЕСТОК / ЧИЖИК / МАГНИТ);
        col 0      — подписи строк (бренды);
        col 13     — пустой столбец-разделитель между годами.

    Датафрейм:
        index   — бренды (в любом регистре), либо в колонке 'index' после reset_index();
        columns — подписи месяцев 'янв.25' ... 'фев.26' (один непрерывный диапазон);
        значения — абсолютные (затраты / TVR) или доли.

    Логика: каждая колонка df разбирается на (месяц, год); по году значение
    направляется в левый ({ly}) или правый ({c_year}) блок, по месяцу — в нужный столбец.
    Годы определяются автоматически (меньший -> {ly}, больший -> {c_year}),
    либо задаются явно через ly_year / c_year (двузначные, например 25 и 26).
    """
    data = df.copy()
    if "index" in data.columns:
        data = data.set_index("index")
    data.index = data.index.map(str)
    data.columns = data.columns.map(str)

    # Если месяцы оказались в индексе, а бренды в колонках — транспонируем.
    cols_are_months = sum(_parse_month_year_label(c)[0] is not None for c in data.columns)
    rows_are_months = sum(_parse_month_year_label(i)[0] is not None for i in data.index)
    if rows_are_months > cols_are_months:
        data = data.T

    # Разбираем колонки df на (месяц, год).
    parsed_cols = {}
    years_seen = set()
    for col in data.columns:
        m, y = _parse_month_year_label(col)
        if m is not None:
            parsed_cols[col] = (m, y)
            years_seen.add(y)

    if not parsed_cols:
        if verbose:
            print(f"  [{name}] SOS/SOV: не удалось распознать месяцы в колонках df")
        return 0

    if ly_year is None or c_year is None:
        ordered_years = sorted(years_seen)
        ly_year = ordered_years[0] if ly_year is None else ly_year
        c_year = ordered_years[-1] if c_year is None else c_year

    # Карты {месяц: ppt_col} для обоих годовых блоков.
    left_cols, right_cols = _detect_year_blocks_from_number_row(
        table, number_row=number_row, label_cols=label_col + 1
    )

    # Карта подписей строк PPT -> имя строки в df.
    df_index_lookup = {_normalize_promo_label(i): i for i in data.index}

    def fmt(value):
        if value_format is not None:
            return value_format(value)
        if as_percent:
            return _format_promo_value(value, as_percent=True, decimals=decimals, suffix=suffix)
        # абсолютные числа: разделитель тысяч пробелом
        if _promo_value_is_empty(value):
            return empty_text
        try:
            number = float(value)
        except (TypeError, ValueError):
            return str(value).strip()
        if number == 0:
            return empty_text
        if abs(number) >= 100:
            return f"{number:,.0f}".replace(",", " ") + suffix
        return f"{number:.{decimals}f}".replace(".", ",") + suffix

    R = len(table.rows)
    written = 0
    unmatched_rows = []

    for r in range(data_start_row, R):
        ppt_label = _normalize_promo_label(table.cell(r, label_col).text)
        df_row = df_index_lookup.get(ppt_label)
        if df_row is None:
            if ppt_label:
                unmatched_rows.append(table.cell(r, label_col).text)
            continue

        for col, (month_num, year2) in parsed_cols.items():
            if year2 == ly_year:
                target = left_cols.get(month_num)
            elif year2 == c_year:
                target = right_cols.get(month_num)
            else:
                target = None

            if target is None:
                continue

            value = data.loc[df_row, col]
            if isinstance(value, pd.Series):
                value = value.dropna()
                value = value.iloc[0] if len(value) else None

            text = empty_text if _promo_value_is_empty(value) else fmt(value)
            _set_ppt_table_cell_text(table.cell(r, target), text, font_size=font_size)
            written += 1

    if verbose:
        print(f"  [{name}] SOS/SOV: записано {written} ячеек "
              f"(ly={ly_year}, c_year={c_year})"
              + (f", не сопоставлены строки: {unmatched_rows}" if unmatched_rows else ""))

    return written


def fill_sos_sov_table_on_slide(
    presentation,
    slide_num,
    df,
    shape_name=None,
    fallback_to_first_table=True,
    **kwargs,
):
    """Обёртка для SOS/SOV-таблиц на слайдах 25/31 (по имени shape)."""
    slide = presentation.slides[slide_num - 1]

    table_shape = None
    if shape_name is not None:
        table_shape = find_shape_by_name(slide, shape_name)
    if table_shape is None and fallback_to_first_table:
        table_shape = find_first_table_on_slide(slide)

    if table_shape is None or not table_shape.has_table:
        raise ValueError(
            f"Не найдена таблица на слайде {slide_num}. shape_name={shape_name!r}"
        )

    fill_sos_sov_table(
        table_shape.table,
        df,
        name=kwargs.pop("name", shape_name or f"slide{slide_num}"),
        **kwargs,
    )
    return table_shape



def update_chart_values_in_place(chart, series_values, categories):
    """
    series_values: dict {имя_серии: [числа]}
    """
    NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    NSA = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    C = NS['c']
    pa = chart._chartSpace.find('.//c:plotArea', NS)
    matched = set() 
    norm_values = { _norm(k): v for k, v in series_values.items()}

    for ser in pa.findall('.//c:ser', NS):
        # 1. достать имя, нормализовав None -> ""
        v = ser.find('c:tx//c:v', NS)
        name = v.text if v is not None and v.text is not None else ""
        # 2. если имени нет в series_values -> continue
        if _norm(name) not in norm_values:
            print(norm_values)
            continue
        new_values = norm_values[_norm(name)]
        matched.add(_norm(name))
        # 3. найти numCache
        cache = ser.find('c:val//c:numCache', NS)
        # 4. удалить старые c:pt
        for pt in list(cache.findall("c:pt", NS)):
            cache.remove(pt)
        # 5. поставить ptCount
        ptcount = cache.find("c:ptCount", NS)
        ptcount.set("val", str(len(new_values)))
        # 6. создать новые c:pt
        for i, value in enumerate(new_values):
            pt = etree.SubElement(cache, f"{{{C}}}pt")
            pt.set("idx", str(i))
            v_el = etree.SubElement(pt, f"{{{C}}}v")
            v_el.text = "0" if value is None else str(value)
        # 7. расширить диапазоны под фактическое число точек
        n = len(new_values)
        for path in ("c:val//c:f", "c:cat//c:f"):
            f_el = ser.find(path, NS)
            if f_el is None:
                continue
            f_el.text = _resize_range(f_el.text, n)

        # 8
        cat_cache = ser.find('c:cat//c:numCache', NS)
        if cat_cache is not None:
            for pt in list(cat_cache.findall('c:pt', NS)):
                cat_cache.remove(pt)
            cat_cache.find('c:ptCount', NS).set('val', str(len(categories)))
            for i, label in enumerate(categories):
                serial = _cat_to_excel_serial(label)
                if serial is None:
                    continue
                pt = etree.SubElement(cat_cache, f'{{{C}}}pt')
                pt.set('idx', str(i))
                v_el = etree.SubElement(pt, f'{{{C}}}v')
                v_el.text = str(serial)

        # 9. дописать dLbl для новых категорий
        dlbls = ser.find('c:dLbls', NS)
        if dlbls is not None:
            existing = dlbls.findall('c:dLbl', NS)
            if existing:
                have = {int(d.find('c:idx', NS).get('val')) for d in existing}
                template_dlbl = existing[-1]          # последний как образец
                for i in range(len(new_values)):
                    if i in have:
                        continue
                    new_d = copy.deepcopy(template_dlbl)
                    new_d.find('c:idx', NS).set('val', str(i))

                    # убрать унаследованные идентификаторы
                    for ext in new_d.findall('.//c:extLst', NS):
                        ext.getparent().remove(ext)
                    for fld in new_d.findall('.//a:fld', NSA):
                        fld.getparent().remove(fld)

                    template_dlbl.addnext(new_d)
                    template_dlbl = new_d


    missing = set(norm_values) - matched
    if missing:
        print(f"ВНИМАНИЕ: серии не найдены в шаблоне: {missing}") 

    
