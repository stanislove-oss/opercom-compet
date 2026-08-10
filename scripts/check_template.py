#!/usr/bin/env python3
"""Сверяет ноутбук и шаблон PowerPoint. Ничего не выполняет и не изменяет.

Ноутбук адресует графики и таблицы по именам объектов (`chart__...`,
`table__...`). Если в шаблоне объект переименовали или удалили, прогон падает —
или, хуже, тихо не заполняет часть слайдов. Эта проверка ловит расхождение за
секунду вместо сорокаминутного прогона.

    python scripts/check_template.py
    python scripts/check_template.py --show-unreferenced
"""

from __future__ import annotations

import argparse
import ast
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from opercom import tags as T  # noqa: E402
from opercom.config import load_settings  # noqa: E402
from opercom.notebook_runner import load_notebook  # noqa: E402

#: Словари ноутбука вида {ключ: {"slide": N, "shape_name": "...", ...}}.
SPEC_DICTS = {
    "CHART_SPECS": "chart",
    "MINI_TABLE_SPECS_6_7": "table",
    "TABLE_SPECS_33_42": "table",
    "PROMO_TABLE_SPECS_S23": "table",
    "SOS_SOV_TABLE_SPECS": "table",
}

#: Имена объектов, которые ноутбук передаёт прямым аргументом (таблицы 13-14).
#: Ровно два подчёркивания и непустой хвост — иначе под шаблон попадают ключи
#: словарей (`chart_type`) и префикс f-строки `table__` из fill_delta_tables.
SHAPE_NAME_RE = re.compile(r"^(?:chart|table)__[A-Za-z0-9_]+$")

#: Правило из ноутбука (`is_delta_table_dataframe_name`): по каким датафреймам
#: fill_delta_tables ищет таблицу с именем `table__{имя_датафрейма}`.
def is_delta_name(name: str) -> bool:
    key = name.lower()
    return "_labels" in key or "_label" in key or "delta" in key


def collect_specs(cells: list[dict]) -> tuple[dict[str, list[dict]], set[str], set[str]]:
    """Достаёт из боевых ячеек словари спецификаций, имена объектов и ключи `dataframes`.

    Читаем через AST, а не выполняем: словари в ноутбуке — литералы, их достаточно
    разобрать статически.
    """
    specs: dict[str, list[dict]] = {}
    literal_names: set[str] = set()
    dataframe_keys: set[str] = set()

    for item in T.select_cells(cells, data_source="db"):
        if not item.should_run:
            continue
        source = "".join(cells[item.index].get("source", []))
        try:
            tree = ast.parse(source)
        except SyntaxError:
            continue

        for node in ast.walk(tree):
            if isinstance(node, ast.Constant) and isinstance(node.value, str):
                if SHAPE_NAME_RE.match(node.value):
                    literal_names.add(node.value)

            # dataframes['ключ'] — так ноутбук объявляет все свои витрины.
            if (
                isinstance(node, ast.Subscript)
                and isinstance(node.value, ast.Name)
                and node.value.id == "dataframes"
                and isinstance(node.slice, ast.Constant)
                and isinstance(node.slice.value, str)
            ):
                dataframe_keys.add(node.slice.value)

            if not isinstance(node, ast.Assign) or not isinstance(node.value, ast.Dict):
                continue
            for target in node.targets:
                if not isinstance(target, ast.Name) or target.id not in SPEC_DICTS:
                    continue
                try:
                    value = ast.literal_eval(node.value)
                except ValueError:
                    continue
                entries = [
                    {"key": key, **entry}
                    for key, entry in value.items()
                    if isinstance(entry, dict) and "shape_name" in entry
                ]
                if entries:
                    specs.setdefault(target.id, []).extend(entries)

    return specs, literal_names, dataframe_keys


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--show-unreferenced", action="store_true", help="показать объекты шаблона без ссылок")
    args = parser.parse_args()

    try:
        import pptx
    except ImportError:
        print("Нужен python-pptx: pip install -r requirements.txt", file=sys.stderr)
        return 1

    settings = load_settings()
    if not settings.template_pptx.is_file():
        print(f"Шаблон не найден: {settings.template_pptx}", file=sys.stderr)
        return 1

    notebook = load_notebook(settings.notebook_path)
    specs, literal_names, dataframe_keys = collect_specs(notebook["cells"])

    presentation = pptx.Presentation(str(settings.template_pptx))
    shapes_by_slide: dict[int, dict[str, object]] = {}
    all_shape_names: set[str] = set()
    for number, slide in enumerate(presentation.slides, start=1):
        shapes_by_slide[number] = {shape.name: shape for shape in slide.shapes}
        all_shape_names.update(shapes_by_slide[number])

    problems: list[str] = []
    checked = 0

    for dict_name, entries in sorted(specs.items()):
        expect_chart = SPEC_DICTS[dict_name] == "chart"
        for entry in entries:
            checked += 1
            slide_num, shape_name, key = entry["slide"], entry["shape_name"], entry["key"]

            if slide_num not in shapes_by_slide:
                problems.append(f"{dict_name}/{key}: в шаблоне нет слайда {slide_num}")
                continue

            shape = shapes_by_slide[slide_num].get(shape_name)
            if shape is None:
                problems.append(
                    f"{dict_name}/{key}: на слайде {slide_num} нет объекта {shape_name!r}"
                )
                continue

            if expect_chart and not shape.has_chart:
                problems.append(f"{dict_name}/{key}: {shape_name!r} на слайде {slide_num} — не график")
            if not expect_chart and not getattr(shape, "has_table", False):
                problems.append(f"{dict_name}/{key}: {shape_name!r} на слайде {slide_num} — не таблица")

        print(f"  {dict_name}: {len(entries)} записей")

    # Имена, переданные прямым аргументом (таблицы длительности на слайдах 13-14).
    spec_names = {entry["shape_name"] for entries in specs.values() for entry in entries}
    for name in sorted(literal_names - spec_names):
        checked += 1
        if name not in all_shape_names:
            problems.append(f"прямая ссылка: объекта {name!r} нет ни на одном слайде")

    print(f"\nПроверено ссылок: {checked}; слайдов в шаблоне: {len(presentation.slides)}")

    # Таблицы дельт: fill_delta_tables ищет `table__{имя_датафрейма}` и молча
    # пропускает промах с любой стороны. Сверяем обе стороны явно.
    delta_shapes = {
        name for name in all_shape_names
        if name.startswith("table__") and is_delta_name(name[len("table__"):])
    }
    delta_dataframes = {key for key in dataframe_keys if is_delta_name(key)}

    orphan_shapes = sorted(delta_shapes - {f"table__{key}" for key in delta_dataframes})
    orphan_dataframes = sorted(
        key for key in delta_dataframes if f"table__{key}" not in all_shape_names
    )

    print(
        f"\nТаблицы дельт: {len(delta_dataframes)} датафреймов, "
        f"{len(delta_shapes)} таблиц в шаблоне"
    )
    if orphan_shapes:
        print("  Таблицы шаблона, для которых нет датафрейма (останутся пустыми):")
        for name in orphan_shapes:
            print(f"    - {name}")
    if orphan_dataframes:
        print("  Датафреймы дельт, для которых нет таблицы в шаблоне (не попадут в презентацию):")
        for key in orphan_dataframes:
            print(f"    - {key}")
    if not orphan_shapes and not orphan_dataframes:
        print("  Совпадают полностью.")

    if args.show_unreferenced:
        unreferenced = sorted(
            name for name in all_shape_names
            if SHAPE_NAME_RE.match(name)
            and name not in spec_names
            and name not in literal_names
            and name not in delta_shapes
        )
        print(f"\nПрочие объекты шаблона без ссылки в ноутбуке ({len(unreferenced)}):")
        for name in unreferenced:
            print(f"  - {name}")

    if problems:
        print(f"\nНайдено расхождений: {len(problems)}", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        return 1

    print("\nOK: все объекты, к которым обращается ноутбук, есть в шаблоне.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
